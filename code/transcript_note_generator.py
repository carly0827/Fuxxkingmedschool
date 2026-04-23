from pathlib import Path
import hashlib
import json
import os
import re
import subprocess
import tempfile
import shutil
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple
import unicodedata

import fitz  # PyMuPDF
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches


@dataclass
class TranscriptBlock:
    timecode: str
    source_name: str
    hinted_page: Optional[int]
    body: str


@dataclass
class LogicalPageRef:
    source: str
    page_no: int
    reason: Optional[str] = None


STOPWORDS = {
    "그리고", "그다음", "이제", "우리가", "여기", "저기", "그냥", "이런", "그런",
    "있는", "하면", "하는", "있고", "있죠", "있어요", "되는", "같아요", "정도",
    "이쪽", "저쪽", "아까", "또", "좀", "더", "때문", "의해", "관련", "이야기",
    "설명", "부분", "내용", "보면", "보세요", "중요", "정리", "가지", "하나",
    "there", "with", "from", "that", "this", "have", "their", "they", "which",
    "page", "pages", "the", "and", "for", "are", "was", "were", "than", "then",
    "upper", "lower", "right", "left", "part", "area", "region", "space", "slide",
}

QUESTION_MARKERS = [
    "25Y", "모아보기", "정답", "문제", "quiz", "question", "객관식", "애매한 문제",
]

KEYWORD_REGEX = re.compile(r"[A-Za-z][A-Za-z0-9_\-]{1,}|[가-힣]{2,}")
TIME_SPLIT_PATTERN = re.compile(
    r"(?ms)^\s*(\d{1,2}:\d{2}(?::\d{2})?)\s*\n+(.+?)\n+·\s*\n+(\d+)페이지\s*\n(.*?)(?=^\s*\d{1,2}:\d{2}(?::\d{2})?\s*$|\Z)"
)

# 핵심어를 의학용어(영문 canonical form)로 맞춘 뒤 비교하기 위한 사전.
# 필요하면 alias를 계속 추가하면 된다.
MEDICAL_ALIAS_GROUPS: Dict[str, List[str]] = {
    "liver": ["liver", "간"],
    "gallbladder": ["gallbladder", "gb", "담낭", "쓸개"],
    "biliary_tree": ["biliary tree", "담도계", "담도"],
    "bile_duct": ["bile duct", "bile ducts", "담관", "담도관"],
    "common_bile_duct": ["common bile duct", "cbd", "총담관"],
    "common_hepatic_duct": ["common hepatic duct", "chd", "총간관"],
    "cystic_duct": ["cystic duct", "담낭관"],
    "hepatic_duct": ["hepatic duct", "간관"],
    "right_hepatic_duct": ["right hepatic duct", "우간관"],
    "left_hepatic_duct": ["left hepatic duct", "좌간관"],
    "portal_vein": ["portal vein", "portal v", "pv", "문맥", "간문맥", "문맥정맥"],
    "hepatic_vein": ["hepatic vein", "hepatic v", "간정맥"],
    "hepatic_artery": ["hepatic artery", "간동맥"],
    "hepatic_artery_proper": ["hepatic artery proper", "proper hepatic artery", "hpa", "고유간동맥"],
    "splenic_vein": ["splenic vein", "비정맥"],
    "splenic_artery": ["splenic artery", "비동맥"],
    "superior_mesenteric_vein": ["superior mesenteric vein", "smv", "상장간막정맥"],
    "superior_mesenteric_artery": ["superior mesenteric artery", "sma", "상장간막동맥"],
    "inferior_vena_cava": ["inferior vena cava", "ivc", "하대정맥"],
    "superior_vena_cava": ["superior vena cava", "svc", "상대정맥"],
    "abdominal_aorta": ["abdominal aorta", "aorta", "복부대동맥"],
    "celiac_trunk": ["celiac trunk", "복강동맥", "celiac axis"],
    "porta_hepatis": ["porta hepatis", "간문", "hepatic portal"],
    "bare_area": ["bare area", "나간부", "간의 나출부"],
    "cantlie_line": ["cantlie line", "캔틀리선"],
    "couinaud_segment": ["couinaud", "segment", "간분절", "세그먼트"],
    "right_lobe": ["right lobe", "우엽"],
    "left_lobe": ["left lobe", "좌엽"],
    "caudate_lobe": ["caudate lobe", "꼬리엽", "미상엽"],
    "quadrate_lobe": ["quadrate lobe", "사각엽"],
    "falciform_ligament": ["falciform ligament", "겸상인대"],
    "coronary_ligament": ["coronary ligament", "관상인대"],
    "triangular_ligament": ["triangular ligament", "삼각인대"],
    "ligamentum_venosum": ["ligamentum venosum", "정맥관인대", "venous ligament"],
    "ligamentum_teres": ["ligamentum teres", "round ligament", "원인대", "원형인대"],
    "hepatoduodenal_ligament": ["hepatoduodenal ligament", "간십이지장인대"],
    "hepatogastric_ligament": ["hepatogastric ligament", "간위인대"],
    "lesser_omentum": ["lesser omentum", "소망", "작은그물막"],
    "greater_omentum": ["greater omentum", "대망", "큰그물막"],
    "lobule": ["lobule", "소엽"],
    "hepatic_lobule": ["hepatic lobule", "간소엽"],
    "portal_lobule": ["portal lobule", "문맥소엽"],
    "acinus": ["acinus", "선방"],
    "hepatic_acinus": ["hepatic acinus", "간선방"],
    "portal_triad": ["portal triad", "문맥삼합", "portal canal"],
    "hepatocyte": ["hepatocyte", "간세포"],
    "sinusoid": ["sinusoid", "sinusoids", "굴모세혈관", "동모양혈관"],
    "kupffer_cell": ["kupffer cell", "쿠퍼세포"],
    "pit_cell": ["pit cell", "pit cells"],
    "stellate_cell": ["stellate cell", "hepatic stellate cell", "성상세포", "ito cell", "ito cells", "이토세포"],
    "space_of_disse": ["space of disse", "disse space", "disse", "디세강"],
    "duodenum": ["duodenum", "십이지장"],
    "stomach": ["stomach", "위"],
    "esophagus": ["esophagus", "oesophagus", "식도"],
    "pancreas": ["pancreas", "췌장", "이자"],
    "pancreatic_duct": ["pancreatic duct", "이자관", "췌관"],
    "main_pancreatic_duct": ["main pancreatic duct", "duct of wirsung", "주이자관", "주췌관", "위르숭관"],
    "accessory_pancreatic_duct": ["accessory pancreatic duct", "duct of santorini", "부이자관", "부췌관", "산토리니관"],
    "ampulla_of_vater": ["ampulla of vater", "vater ampulla", "바터팽대", "간췌팽대"],
    "sphincter_of_oddi": ["sphincter of oddi", "oddi sphincter", "오디조임근", "간췌팽대조임근"],
    "major_duodenal_papilla": ["major duodenal papilla", "주유두", "큰십이지장유두"],
    "minor_duodenal_papilla": ["minor duodenal papilla", "부유두", "작은십이지장유두"],
    "choledochoduodenal_junction": ["choledochoduodenal junction", "총담관십이지장연접부"],
    "hepatocystic_triangle": ["hepatocystic triangle", "triangle of calot", "간담낭삼각", "칼로삼각"],
    "luschka_duct": ["duct of luschka", "luschka duct", "루슈카관"],
    "rokitansky_aschoff_sinus": ["rokitansky aschoff sinus", "rokitansky-aschoff sinus", "로키탄스키 아쇼프굴", "로키탄스키아쇼프굴"],
    "spleen": ["spleen", "비장"],
    "splenic_hilum": ["splenic hilum", "splenic hilus", "비문"],
    "gastrosplenic_ligament": ["gastrosplenic ligament", "위비인대"],
    "splenorenal_ligament": ["splenorenal ligament", "lienorenal ligament", "비신인대"],
    "phrenicocolic_ligament": ["phrenicocolic ligament", "횡격막결장인대"],
    "mucosa": ["mucosa", "점막"],
    "submucosa": ["submucosa", "점막하층"],
    "muscularis": ["muscularis", "muscle layer", "근육층"],
    "serosa": ["serosa", "장막"],
    "adventitia": ["adventitia", "외막"],
    "lamina_propria": ["lamina propria", "고유판"],
    "simple_columnar_epithelium": ["simple columnar epithelium", "단층원주상피"],
    "microvilli": ["microvilli", "미세융모"],
    "smooth_muscle": ["smooth muscle", "평활근"],
}

MEDICAL_ALIAS_PATTERNS: List[Tuple[re.Pattern, str]] = []
for canon, aliases in MEDICAL_ALIAS_GROUPS.items():
    seen = set()
    candidates = [canon.replace("_", " "), canon] + list(aliases)
    for alias in sorted(candidates, key=len, reverse=True):
        alias_low = unicodedata.normalize("NFC", alias.lower().strip())
        if not alias_low or alias_low in seen:
            continue
        seen.add(alias_low)
        if re.search(r"[a-z]", alias_low):
            pattern = re.compile(r"(?<![a-z])" + re.escape(alias_low) + r"(?![a-z])")
        else:
            pattern = re.compile(re.escape(alias_low))
        MEDICAL_ALIAS_PATTERNS.append((pattern, canon))
MEDICAL_ALIAS_PATTERNS.sort(key=lambda x: len(x[0].pattern), reverse=True)


def find_korean_font() -> Optional[str]:
    candidates = [
        "Noto Sans CJK KR",
        "Noto Sans CJK JP",
        "NanumGothic",
        "NanumBarunGothic",
        "UnDotum",
        "Malgun Gothic",
        "Apple SD Gothic Neo",
    ]
    for name in candidates:
        try:
            out = subprocess.check_output(["fc-match", "-f", "%{file}\n", name], text=True).strip()
            if out and os.path.exists(out):
                return out
        except Exception:
            continue
    fallback_paths = [
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "C:/Windows/Fonts/malgun.ttf",
    ]
    for path in fallback_paths:
        if os.path.exists(path):
            return path
    return None


def normalize_medical_terms(text: str) -> str:
    text = unicodedata.normalize("NFC", text or "")
    text = text.lower()
    text = text.replace("\r\n", "\n")
    for pattern, canon in MEDICAL_ALIAS_PATTERNS:
        text = pattern.sub(canon, text)
    text = re.sub(r"[^0-9a-z가-힣_\-\s]", " ", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_text(text: str) -> str:
    text = normalize_medical_terms(text)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_for_exact_compare(text: str) -> str:
    text = text.replace("\r\n", "\n")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def sha1_text(text: str) -> str:
    return hashlib.sha1(text.encode("utf-8")).hexdigest()


def tokenize(text: str) -> List[str]:
    normalized = normalize_medical_terms(text)
    return [tok for tok in KEYWORD_REGEX.findall(normalized) if len(tok) >= 2]


def extract_keywords(text: str, limit: int = 8) -> List[str]:
    counts: Dict[str, int] = {}
    original: Dict[str, str] = {}
    for tok in tokenize(text):
        low = tok.lower()
        if low in STOPWORDS:
            continue
        counts[low] = counts.get(low, 0) + 1
        original.setdefault(low, tok)
    ranked = sorted(counts.items(), key=lambda kv: (-kv[1], -len(kv[0]), kv[0]))
    return [original[low] for low, _ in ranked[:limit]]



def _color_to_rgb255(color) -> Optional[Tuple[int, int, int]]:
    if color is None:
        return None
    if isinstance(color, (list, tuple)) and len(color) >= 3:
        vals = list(color[:3])
        if vals and max(vals) <= 1.0:
            return tuple(max(0, min(255, int(round(v * 255)))) for v in vals[:3])
        return tuple(max(0, min(255, int(round(v)))) for v in vals[:3])
    if isinstance(color, int):
        return ((color >> 16) & 255, (color >> 8) & 255, color & 255)
    return None


TARGET_GREEN_HIGHLIGHT_RGB = (0xA5, 0xC8, 0x91)  # #A5C891
TARGET_GREEN_HANDWRITING_RGB = (0x00, 0x73, 0x55)  # #007355
HIGHLIGHT_COLOR_TOLERANCE = 64
HANDWRITING_COLOR_TOLERANCE = 42

TARGET_RED_EMPHASIS_RGB = (240, 176, 176)
TARGET_RED_HANDWRITING_RGB = (214, 68, 68)
RED_EMPHASIS_TOLERANCE = 88
RED_HANDWRITING_TOLERANCE = 84


def _color_distance(rgb_a: Tuple[int, int, int], rgb_b: Tuple[int, int, int]) -> float:
    return ((rgb_a[0] - rgb_b[0]) ** 2 + (rgb_a[1] - rgb_b[1]) ** 2 + (rgb_a[2] - rgb_b[2]) ** 2) ** 0.5



def _matches_specific_green(color, *, mode: str = "either") -> bool:
    """
    사용자가 지정한 두 색만 초록 우선 단서로 인정한다.
    - 형광펜: #A5C891
    - 필기색:  #007355
    PDF 렌더/annotation 저장 과정에서 약간의 색 흔들림이 있어 소폭 tolerance를 둔다.
    """
    rgb = _color_to_rgb255(color)
    if not rgb:
        return False

    mode = (mode or "either").lower()
    if mode == "highlight":
        return _color_distance(rgb, TARGET_GREEN_HIGHLIGHT_RGB) <= HIGHLIGHT_COLOR_TOLERANCE
    if mode == "handwriting":
        return _color_distance(rgb, TARGET_GREEN_HANDWRITING_RGB) <= HANDWRITING_COLOR_TOLERANCE
    return (
        _color_distance(rgb, TARGET_GREEN_HIGHLIGHT_RGB) <= HIGHLIGHT_COLOR_TOLERANCE
        or _color_distance(rgb, TARGET_GREEN_HANDWRITING_RGB) <= HANDWRITING_COLOR_TOLERANCE
    )


def _matches_red_emphasis(color, *, mode: str = "either") -> bool:
    """
    붉은색 표시(형광펜/밑줄/필기) 단서를 잡는다.
    사용자가 예시로 준 붉은 표시 이미지를 기준으로, 붉은 분홍 계열 highlight와 진한 붉은 필기 모두 허용한다.
    PDF 저장 시 색이 약간 흔들릴 수 있어 근접색도 함께 허용한다.
    """
    rgb = _color_to_rgb255(color)
    if not rgb:
        return False
    r, g, b = rgb
    highlight_like = _color_distance(rgb, TARGET_RED_EMPHASIS_RGB) <= RED_EMPHASIS_TOLERANCE or (r >= 170 and r > g + 18 and r > b + 8)
    handwriting_like = _color_distance(rgb, TARGET_RED_HANDWRITING_RGB) <= RED_HANDWRITING_TOLERANCE or (r >= 110 and r > g + 35 and r > b + 20)
    mode = (mode or "either").lower()
    if mode == "highlight":
        return highlight_like
    if mode == "handwriting":
        return handwriting_like
    return highlight_like or handwriting_like


def _dedupe_keep_order(items: List[str]) -> List[str]:
    seen = set()
    out: List[str] = []
    for item in items:
        normalized = re.sub(r"\s+", " ", item).strip()
        if not normalized:
            continue
        key = normalized.lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(normalized)
    return out


def _rect_from_drawing(draw_item: Dict[str, object]) -> Optional[fitz.Rect]:
    rect = draw_item.get("rect")
    if isinstance(rect, fitz.Rect):
        return rect
    items = draw_item.get("items", [])
    xs: List[float] = []
    ys: List[float] = []
    for item in items:
        for part in item[1:]:
            if isinstance(part, fitz.Point):
                xs.append(part.x)
                ys.append(part.y)
            elif isinstance(part, fitz.Rect):
                xs.extend([part.x0, part.x1])
                ys.extend([part.y0, part.y1])
            elif isinstance(part, (list, tuple)) and len(part) >= 2:
                try:
                    xs.append(float(part[0]))
                    ys.append(float(part[1]))
                except Exception:
                    pass
    if xs and ys:
        return fitz.Rect(min(xs), min(ys), max(xs), max(ys))
    return None


def extract_page_color_priority(page: fitz.Page) -> Dict[str, object]:
    """
    사용자가 지정한 색 단서를 우선 추출한다.
    - 초록 형광펜: #A5C891
    - 초록 필기:   #007355
    - 붉은 강조:   예시 이미지에 맞춘 붉은/분홍 강조와 붉은 필기 근접색

    반환값은 전사문 페이지 부착 매칭의 1순위(초록)와,
    정리본 보강용 붉은 강조 단서를 함께 담는다.
    """
    green_text_spans: List[str] = []
    green_highlight_texts: List[str] = []
    red_text_spans: List[str] = []
    red_highlight_texts: List[str] = []

    try:
        text_dict = page.get_text("dict")
    except Exception:
        text_dict = {"blocks": []}

    for block in text_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            line_parts: List[str] = []
            line_is_green = False
            line_is_red = False
            for span in line.get("spans", []):
                span_text = re.sub(r"\s+", " ", span.get("text", "")).strip()
                if span_text:
                    line_parts.append(span_text)
                if span_text and _matches_specific_green(span.get("color"), mode="handwriting"):
                    line_is_green = True
                    green_text_spans.append(span_text)
                if span_text and _matches_red_emphasis(span.get("color"), mode="handwriting"):
                    line_is_red = True
                    red_text_spans.append(span_text)
            if line_is_green and line_parts:
                green_text_spans.append(" ".join(line_parts))
            if line_is_red and line_parts:
                red_text_spans.append(" ".join(line_parts))

    try:
        annot = page.first_annot
        while annot:
            subtype = ""
            try:
                subtype = (annot.type[1] or "").lower()
            except Exception:
                subtype = ""
            colors = getattr(annot, "colors", {}) or {}
            stroke = colors.get("stroke")
            fill = colors.get("fill")
            is_markup = ("highlight" in subtype or "squiggly" in subtype or "underline" in subtype)
            if is_markup and (_matches_specific_green(stroke, mode="highlight") or _matches_specific_green(fill, mode="highlight") or _matches_specific_green(stroke, mode="handwriting") or _matches_specific_green(fill, mode="handwriting")):
                base_rect = annot.rect
                probe = fitz.Rect(base_rect.x0 - 3, base_rect.y0 - 2, base_rect.x1 + 3, base_rect.y1 + 2)
                extracted = re.sub(r"\s+", " ", page.get_textbox(probe) or "").strip()
                if extracted:
                    green_highlight_texts.append(extracted)
            if is_markup and (_matches_red_emphasis(stroke, mode="highlight") or _matches_red_emphasis(fill, mode="highlight") or _matches_red_emphasis(stroke, mode="handwriting") or _matches_red_emphasis(fill, mode="handwriting")):
                base_rect = annot.rect
                probe = fitz.Rect(base_rect.x0 - 3, base_rect.y0 - 2, base_rect.x1 + 3, base_rect.y1 + 2)
                extracted = re.sub(r"\s+", " ", page.get_textbox(probe) or "").strip()
                if extracted:
                    red_highlight_texts.append(extracted)
            annot = annot.next
    except Exception:
        pass

    try:
        for draw_item in page.get_drawings():
            fill = draw_item.get("fill")
            stroke = draw_item.get("color") or draw_item.get("stroke")
            rect = _rect_from_drawing(draw_item)
            if not rect:
                continue
            if rect.width < 18 or rect.height < 4 or rect.height > 90:
                continue
            probe = fitz.Rect(rect.x0 - 2, rect.y0 - 2, rect.x1 + 2, rect.y1 + 2)
            extracted = re.sub(r"\s+", " ", page.get_textbox(probe) or "").strip()
            if (_matches_specific_green(fill, mode="highlight") or _matches_specific_green(fill, mode="handwriting") or _matches_specific_green(stroke, mode="highlight") or _matches_specific_green(stroke, mode="handwriting")) and extracted:
                green_highlight_texts.append(extracted)
            if (_matches_red_emphasis(fill, mode="highlight") or _matches_red_emphasis(fill, mode="handwriting") or _matches_red_emphasis(stroke, mode="highlight") or _matches_red_emphasis(stroke, mode="handwriting")) and extracted:
                red_highlight_texts.append(extracted)
    except Exception:
        pass

    green_text_spans = _dedupe_keep_order(green_text_spans)
    green_highlight_texts = _dedupe_keep_order(green_highlight_texts)
    red_text_spans = _dedupe_keep_order(red_text_spans)
    red_highlight_texts = _dedupe_keep_order(red_highlight_texts)
    priority_fragments = _dedupe_keep_order(green_highlight_texts + green_text_spans)
    priority_text = "\n".join(priority_fragments)
    red_fragments = _dedupe_keep_order(red_highlight_texts + red_text_spans)
    red_text = "\n".join(red_fragments)

    return {
        "green_text_spans": green_text_spans,
        "green_highlight_texts": green_highlight_texts,
        "priority_fragments": priority_fragments,
        "priority_text": priority_text,
        "priority_keywords": extract_keywords(priority_text, limit=16),
        "green_highlight_keywords": extract_keywords(" ".join(green_highlight_texts), limit=10),
        "green_text_keywords": extract_keywords(" ".join(green_text_spans), limit=10),
        "red_text_spans": red_text_spans,
        "red_highlight_texts": red_highlight_texts,
        "red_fragments": red_fragments,
        "red_text": red_text,
        "red_keywords": extract_keywords(red_text, limit=12),
    }

def build_page_term_hints(page_text: str) -> List[str]:
    lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]
    if not lines:
        return []
    top_lines = []
    for ln in lines[:8]:
        cleaned = re.sub(r"^[•▪\-\d\)\.\s]+", "", ln).strip()
        if cleaned:
            top_lines.append(cleaned)
    return extract_keywords(" ".join(top_lines), limit=12)


def summarize_transcript_for_matching(body: str, keyword_limit: int = 8) -> str:
    body = re.sub(r"\s+", " ", body).strip()
    if not body:
        return ""
    first = re.split(r"(?<=[.!?])\s+|(?<=다)\s+", body)[0].strip()
    keywords = extract_keywords(body, limit=keyword_limit)
    parts = [first]
    if keywords:
        parts.append("의학핵심어: " + ", ".join(keywords))
    return "\n".join(parts).strip()


def parse_transcript(transcript_text: str) -> List[TranscriptBlock]:
    text = transcript_text.replace("\r\n", "\n").strip()
    blocks: List[TranscriptBlock] = []
    for m in TIME_SPLIT_PATTERN.finditer(text):
        timecode, source_name, page_s, body = m.groups()
        body = re.sub(r"\s+", " ", body).strip()
        blocks.append(TranscriptBlock(timecode.strip(), source_name.strip(), int(page_s), body))
    if blocks:
        return blocks

    chunks = re.split(r"(?m)^\s*(?=\d{1,2}:\d{2}(?::\d{2})?\s*$)", text)
    for chunk in chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
        if len(lines) < 4:
            continue
        timecode = lines[0]
        hinted_page = None
        source_name = ""
        body_start = 0
        for i, ln in enumerate(lines[1:], start=1):
            pg = re.search(r"(\d+)페이지", ln)
            if pg:
                hinted_page = int(pg.group(1))
                body_start = i + 1
                break
            if ln != "·" and not source_name:
                source_name = ln
        if hinted_page is None:
            continue
        body = " ".join(lines[body_start:]).strip()
        blocks.append(TranscriptBlock(timecode, source_name, hinted_page, body))
    return blocks


def classify_page(page_text: str, page_index: int) -> str:
    text = normalize_text(page_text)
    if not text:
        return "blank"
    if any(marker.lower() in text for marker in [m.lower() for m in QUESTION_MARKERS]):
        if "모아보기" in page_text or "25Y" in page_text or "정답" in page_text or "애매한 문제" in page_text:
            return "question"
    lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]
    if len(lines) <= 4 and any(m.lower() in text for m in ["25y", "모아보기"]):
        return "question"
    return "content"


def content_page_sequence_map(page_classes: List[str]) -> Tuple[Dict[int, int], Dict[int, int], List[int]]:
    content_indices = [i for i, cls in enumerate(page_classes) if cls == "content"]
    seq_to_page = {seq + 1: idx for seq, idx in enumerate(content_indices)}
    page_to_seq = {idx: seq + 1 for seq, idx in enumerate(content_indices)}
    return seq_to_page, page_to_seq, content_indices

def is_excluded_from_transcript_page_count(ref: LogicalPageRef, page_class: str, logical_index: int) -> bool:
    """전사문의 페이지 수 힌트는 줄 기준으로 세되, 병합 시 앞에 추가된 야붙 1페이지와 문제만 있는 페이지는 제외한다."""
    if page_class == "question":
        return True
    if logical_index == 0 and ref.source == "yabout" and ref.page_no == 1:
        return True
    if getattr(ref, "reason", None) == "inserted_yabout_front":
        return True
    return False


def transcript_anchor_sequence_map(
    page_plan: List[LogicalPageRef],
    page_classes: List[str],
) -> Tuple[Dict[int, int], Dict[int, int], List[int]]:
    """전사문의 n페이지를 해석할 때 사용할 논리 페이지 순번.

    규칙:
    - 줄+야붙 병합본에서는 앞에 추가된 야붙 1페이지는 세지 않는다.
    - 문제만 있는 페이지(question)는 세지 않는다.
    - 그 외 페이지는 content page 순번으로 센다.
    """
    counted_indices: List[int] = []
    for i, (ref, cls) in enumerate(zip(page_plan, page_classes)):
        if cls != "content":
            continue
        if is_excluded_from_transcript_page_count(ref, cls, i):
            continue
        counted_indices.append(i)
    seq_to_page = {seq + 1: idx for seq, idx in enumerate(counted_indices)}
    page_to_seq = {idx: seq + 1 for seq, idx in enumerate(counted_indices)}
    return seq_to_page, page_to_seq, counted_indices


def equalized_overlap_score(a_keywords: List[str], b_keywords: List[str]) -> float:
    if not a_keywords or not b_keywords:
        return 0.0
    k = min(len(a_keywords), len(b_keywords))
    aset = {kw.lower() for kw in a_keywords[:k]}
    bset = {kw.lower() for kw in b_keywords[:k]}
    if not aset or not bset:
        return 0.0
    return len(aset & bset) / float(k)



def build_page_features(page_text: str, color_priority: Optional[Dict[str, object]] = None) -> Dict[str, object]:
    color_priority = color_priority or {}
    lines = [ln.strip() for ln in page_text.splitlines() if ln.strip()]

    title = ""
    subtitle = ""
    for ln in lines[:6]:
        cleaned = re.sub(r"^[•▪\-\d\)\.\s]+", "", ln).strip()
        if not cleaned:
            continue
        if not title and len(cleaned) <= 100:
            title = cleaned
            continue
        if title and not subtitle and len(cleaned) <= 120:
            subtitle = cleaned
            break

    bullet_candidates = []
    for ln in lines[:18]:
        cleaned = re.sub(r"^[•▪\-\d\)\.\s]+", "", ln).strip()
        if not cleaned or len(cleaned) <= 1:
            continue
        if cleaned == title or cleaned == subtitle:
            continue
        bullet_candidates.append(cleaned)
    core_bullets = bullet_candidates[:5]

    headline_parts = [p for p in [title, subtitle] if p]
    headline_text = " ".join(headline_parts)
    page_scan_text = " ".join(lines)
    priority_fragments = color_priority.get("priority_fragments", [])
    priority_text = color_priority.get("priority_text", "")
    priority_prefixed_scan = (priority_text + "\n" + page_scan_text).strip()

    page_scan_keywords = extract_keywords(priority_prefixed_scan, limit=16)
    title_keywords = extract_keywords(headline_text, limit=8)
    body_keywords = extract_keywords(priority_prefixed_scan, limit=16)
    core_text = " ".join(priority_fragments[:4] + headline_parts + core_bullets[:4] + page_scan_keywords[:8])
    core_keywords = extract_keywords((priority_text + "\n" + core_text).strip(), limit=14)
    page_term_hints = _dedupe_keep_order(color_priority.get("priority_keywords", []) + build_page_term_hints(page_text))[:16]

    summary_parts = []
    if priority_fragments:
        summary_parts.append("[초록 우선 핵심] " + " | ".join(priority_fragments[:3]))
    if title:
        summary_parts.append(title)
    if subtitle and subtitle != title:
        summary_parts.append(subtitle)
    if core_bullets:
        summary_parts.append(" | ".join(core_bullets[:3]))
    if core_keywords:
        summary_parts.append("페이지 의학핵심어: " + ", ".join(core_keywords[:10]))
    summary = "\n".join(summary_parts).strip()

    short_slide = len(tokenize(core_text)) <= 22
    return {
        "title": title,
        "subtitle": subtitle,
        "title_keywords": title_keywords,
        "priority_keywords": color_priority.get("priority_keywords", []),
        "body_keywords": body_keywords,
        "page_scan_keywords": page_scan_keywords,
        "page_term_hints": page_term_hints,
        "summary": summary,
        "summary_keywords": core_keywords[:12],
        "short_slide": short_slide,
        "core_bullets": core_bullets,
        "headline_text": headline_text,
        "normalized_page_text": normalize_medical_terms(priority_prefixed_scan),
        "priority_fragments": priority_fragments,
        "green_text_spans": color_priority.get("green_text_spans", []),
        "green_highlight_texts": color_priority.get("green_highlight_texts", []),
        "red_text_spans": color_priority.get("red_text_spans", []),
        "red_highlight_texts": color_priority.get("red_highlight_texts", []),
        "red_fragments": color_priority.get("red_fragments", []),
        "red_keywords": color_priority.get("red_keywords", []),
    }

def build_block_features(block: TranscriptBlock) -> Dict[str, object]:
    summary = summarize_transcript_for_matching(block.body, keyword_limit=10)
    summary_keywords = extract_keywords(summary, limit=10)
    body_keywords = extract_keywords(block.body, limit=16)
    first_sentence = re.split(r"(?<=[.!?])\s+|(?<=다)\s+", block.body.strip())[0].strip() if block.body.strip() else ""
    medical_keywords = []
    seen = set()
    for kw in summary_keywords + body_keywords:
        if kw not in seen:
            medical_keywords.append(kw)
            seen.add(kw)
    return {
        "summary": summary,
        "summary_keywords": summary_keywords,
        "body_keywords": body_keywords,
        "first_sentence": first_sentence,
        "normalized_body": normalize_medical_terms(block.body),
        "medical_keywords": medical_keywords[:20],
    }



def page_block_score(page_feat: Dict[str, object], block_feat: Dict[str, object]) -> float:
    score = 0.0
    # 1순위: 초록색 필기/형광펜 기반 핵심어
    score += 7.0 * equalized_overlap_score(page_feat.get("priority_keywords", []), block_feat["summary_keywords"])
    score += 4.6 * equalized_overlap_score(page_feat.get("priority_keywords", []), block_feat["body_keywords"])

    # 2~3순위: 기존 제목/부제/본문/의학용어 힌트
    score += 5.5 * equalized_overlap_score(page_feat["summary_keywords"], block_feat["summary_keywords"])
    score += 3.0 * equalized_overlap_score(page_feat["title_keywords"], block_feat["summary_keywords"])
    score += 2.6 * equalized_overlap_score(page_feat["page_term_hints"], block_feat["summary_keywords"])
    score += 2.3 * equalized_overlap_score(page_feat["body_keywords"], block_feat["summary_keywords"])
    score += 1.5 * equalized_overlap_score(page_feat["body_keywords"], block_feat["body_keywords"])

    title_set = {x.lower() for x in page_feat["title_keywords"]}
    subtitle_set = {x.lower() for x in extract_keywords(page_feat.get("subtitle", ""), limit=6)}
    hint_set = {x.lower() for x in page_feat.get("page_term_hints", [])}
    page_scan_set = {x.lower() for x in page_feat.get("page_scan_keywords", [])}
    priority_set = {x.lower() for x in page_feat.get("priority_keywords", [])}
    block_summary_set = {x.lower() for x in block_feat["summary_keywords"]}
    block_body_set = {x.lower() for x in block_feat["body_keywords"]}

    priority_exact = len(priority_set & block_summary_set)
    priority_body_exact = len(priority_set & block_body_set)
    title_exact = len(title_set & block_summary_set)
    title_body_exact = len(title_set & block_body_set)
    subtitle_hits = len(subtitle_set & (block_summary_set | block_body_set))
    hint_hits = len(hint_set & block_body_set)
    scan_hits = len(page_scan_set & block_body_set)

    score += 5.4 * priority_exact
    score += 3.2 * priority_body_exact
    score += 3.6 * title_exact
    score += 2.3 * title_body_exact
    score += 2.0 * subtitle_hits
    score += 1.9 * hint_hits
    score += 1.4 * scan_hits

    core_set = {x.lower() for x in (page_feat["summary_keywords"] + page_feat["title_keywords"] + page_feat.get("page_term_hints", []) + page_feat.get("priority_keywords", []))}
    exact = len(core_set & block_body_set)
    score += 1.9 * exact

    if page_feat["short_slide"] and (priority_exact > 0 or title_exact > 0 or title_body_exact > 0 or subtitle_hits > 0):
        score += 1.6 * max(priority_exact, title_exact, title_body_exact, subtitle_hits)

    first_set = {x.lower() for x in extract_keywords(block_feat["first_sentence"], limit=8)}
    score += 2.2 * len(first_set & priority_set)
    score += 1.4 * len(first_set & (title_set | subtitle_set | hint_set))
    return score



def align_blocks_to_pages(
    blocks: List[TranscriptBlock],
    content_indices: List[int],
    page_features: Dict[int, Dict[str, object]],
    page_to_seq: Dict[int, int],
    transcript_anchor_seq_to_page: Dict[int, int],
) -> Tuple[Dict[int, List[int]], Dict[str, object]]:
    n = len(blocks)
    m = len(content_indices)
    if n == 0 or m == 0:
        return {idx: [] for idx in content_indices}, {"mapping_mode": "empty"}

    block_features = [build_block_features(b) for b in blocks]
    score_matrix = [[-1e9] * m for _ in range(n)]
    hint_windows: List[Tuple[int, int]] = []

    for j, block in enumerate(blocks):
        has_explicit_hint = bool(block.hinted_page and block.hinted_page in transcript_anchor_seq_to_page)
        if has_explicit_hint:
            anchor_page = transcript_anchor_seq_to_page[block.hinted_page]
            anchor_seq = page_to_seq[anchor_page]
        else:
            anchor_seq = min(j + 1, m)

        # 우선순위:
        # 1) 초록색 필기/형광펜 기반 priority score
        # 2) 전사본에 적힌 페이지 번호(anchor_seq)
        # 3) 제목/부제/핵심 bullet/의학용어 힌트/본문 키워드
        window_radius = 1 if has_explicit_hint else 2
        win_lo = max(1, anchor_seq - window_radius)
        win_hi = min(m, anchor_seq + window_radius)
        hint_windows.append((win_lo, win_hi))

        for i, page_idx in enumerate(content_indices, start=1):
            base = page_block_score(page_features[page_idx], block_features[j])
            dist = abs(i - anchor_seq)

            if has_explicit_hint:
                if dist == 0:
                    hint_bonus = 5.6
                elif dist == 1:
                    hint_bonus = 3.0
                elif dist == 2:
                    hint_bonus = 0.7
                else:
                    hint_bonus = -4.0 - 2.3 * (dist - 2)
                local_bonus = 1.5 if win_lo <= i <= win_hi else 0.0
            else:
                if dist == 0:
                    hint_bonus = 3.5
                elif dist == 1:
                    hint_bonus = 1.8
                elif dist == 2:
                    hint_bonus = 0.5
                else:
                    hint_bonus = -2.4 - 1.8 * (dist - 2)
                local_bonus = 0.7 if win_lo <= i <= win_hi else 0.0

            score_matrix[j][i - 1] = base + hint_bonus + local_bonus

    neg = -1e18
    dp = [[neg] * m for _ in range(n)]
    prev = [[-1] * m for _ in range(n)]

    for i in range(m):
        seq_i = i + 1
        dp[0][i] = score_matrix[0][i] - 0.8 * max(0, seq_i - 1)

    for j in range(1, n):
        for i in range(m):
            best = neg
            best_p = -1
            for p in range(i + 1):
                jump = i - p
                penalty = 0.0
                if jump > 2:
                    penalty += 1.0 * (jump - 2)
                if jump == 0:
                    penalty += 0.15
                cand = dp[j - 1][p] + score_matrix[j][i] - penalty
                if cand > best:
                    best = cand
                    best_p = p
            dp[j][i] = best
            prev[j][i] = best_p

    end_i = max(range(m), key=lambda i: dp[n - 1][i])
    assignment = [-1] * n
    cur = end_i
    for j in range(n - 1, -1, -1):
        assignment[j] = cur
        cur = prev[j][cur] if j > 0 else -1

    page_to_block_idxs: Dict[int, List[int]] = {idx: [] for idx in content_indices}
    for j, ai in enumerate(assignment):
        page_to_block_idxs[content_indices[ai]].append(j)

    metrics = {
        "mapping_mode": "monotonic_dp_green_priority_then_transcript_page_anchor_then_title_subtitle_term_hints",
        "assignment_content_seq": [a + 1 for a in assignment],
        "hint_windows": hint_windows,
        "transcript_page_anchor_rule": "줄 기준 페이지수 해석: 줄+야붙 병합 시 앞에 추가된 야붙 1페이지와 문제만 있는 페이지는 제외. 초록 우선 단서는 #007355 필기색과 #A5C891 형광펜색만 반영",
        "transcript_anchor_sequence_pages": sorted(transcript_anchor_seq_to_page.keys()),
    }
    return page_to_block_idxs, metrics



def score_block_to_content_seq(
    block: TranscriptBlock,
    block_feat: Dict[str, object],
    seq_i: int,
    content_indices: List[int],
    page_features: Dict[int, Dict[str, object]],
    page_to_seq: Dict[int, int],
    transcript_anchor_seq_to_page: Dict[int, int],
) -> float:
    page_idx = content_indices[seq_i - 1]
    has_explicit_hint = bool(block.hinted_page and block.hinted_page in transcript_anchor_seq_to_page)
    if has_explicit_hint:
        anchor_page = transcript_anchor_seq_to_page[block.hinted_page]
        anchor_seq = page_to_seq[anchor_page]
    else:
        anchor_seq = min(seq_i, len(content_indices))
    window_radius = 1 if has_explicit_hint else 2
    win_lo = max(1, anchor_seq - window_radius)
    win_hi = min(len(content_indices), anchor_seq + window_radius)
    base = page_block_score(page_features[page_idx], block_feat)
    dist = abs(seq_i - anchor_seq)
    if has_explicit_hint:
        if dist == 0:
            hint_bonus = 5.6
        elif dist == 1:
            hint_bonus = 3.0
        elif dist == 2:
            hint_bonus = 0.7
        else:
            hint_bonus = -4.0 - 2.3 * (dist - 2)
        local_bonus = 1.5 if win_lo <= seq_i <= win_hi else 0.0
    else:
        if dist == 0:
            hint_bonus = 3.5
        elif dist == 1:
            hint_bonus = 1.8
        elif dist == 2:
            hint_bonus = 0.5
        else:
            hint_bonus = -2.4 - 1.8 * (dist - 2)
        local_bonus = 0.7 if win_lo <= seq_i <= win_hi else 0.0
    return base + hint_bonus + local_bonus


def build_assignment_diagnostics(
    blocks: List[TranscriptBlock],
    block_features: List[Dict[str, object]],
    content_indices: List[int],
    page_features: Dict[int, Dict[str, object]],
    page_to_seq: Dict[int, int],
    seq_to_page: Dict[int, int],
    transcript_anchor_seq_to_page: Dict[int, int],
    assignment_content_seq: List[int],
) -> Tuple[List[Dict[str, object]], Dict[str, object]]:
    diagnostics: List[Dict[str, object]] = []

    for j, block in enumerate(blocks):
        assigned_seq = assignment_content_seq[j]
        assigned_page_idx = seq_to_page[assigned_seq]
        row = []
        for seq_i in range(1, len(content_indices) + 1):
            row.append((seq_i, score_block_to_content_seq(block, block_features[j], seq_i, content_indices, page_features, page_to_seq, transcript_anchor_seq_to_page)))
        ranked = sorted(row, key=lambda x: x[1], reverse=True)
        best_seq, best_score = ranked[0]
        second_seq, second_score = ranked[1] if len(ranked) > 1 else (ranked[0][0], ranked[0][1])
        assigned_score = next(score for seq_i, score in ranked if seq_i == assigned_seq)
        margin_vs_best = assigned_score - best_score
        margin_vs_second = assigned_score - second_score if assigned_seq == best_seq else assigned_score - best_score

        title_set = {x.lower() for x in page_features[assigned_page_idx]["title_keywords"]}
        priority_set = {x.lower() for x in page_features[assigned_page_idx].get("priority_keywords", [])}
        body_set = {x.lower() for x in block_features[j]["body_keywords"]}
        summary_set = {x.lower() for x in block_features[j]["summary_keywords"]}
        priority_hits = len(priority_set & (body_set | summary_set))
        title_hits = len(title_set & (body_set | summary_set))
        core_hits = len({x.lower() for x in page_features[assigned_page_idx]["summary_keywords"]} & body_set)

        reasons: List[str] = []
        if block.hinted_page is not None:
            hinted_exists = block.hinted_page in transcript_anchor_seq_to_page
            if hinted_exists and abs(assigned_seq - block.hinted_page) > 2:
                reasons.append("전사문 페이지 힌트와 실제 배정 페이지 거리가 큼")
            if not hinted_exists:
                reasons.append("전사문 페이지 힌트가 '추가된 첫페이지·문제페이지 제외' 기준 페이지수와 직접 대응되지 않음")
        if priority_hits == 0:
            reasons.append("초록색 필기/형광펜 기반 핵심어와 전사문 직접 일치가 약함")
        if title_hits == 0:
            reasons.append("슬라이드 제목과 전사문 핵심어 직접 일치가 약함")
        if core_hits == 0:
            reasons.append("슬라이드 핵심어와 전사문 본문 겹침이 적음")
        if margin_vs_second < 0.6:
            reasons.append("인접 후보 페이지와 점수 차가 작음")
        if assigned_seq != best_seq:
            reasons.append("최고점 페이지가 아닌 위치에 단조 정렬 제약으로 배정됨")

        if assigned_score >= 14.0 and (priority_hits > 0 or title_hits > 0 or margin_vs_second >= 1.5):
            confidence = "high"
        elif assigned_score >= 9.0 and (priority_hits > 0 or margin_vs_second >= 0.5):
            confidence = "medium"
        else:
            confidence = "low"

        top_candidates = []
        for seq_i, score in ranked[:5]:
            page_idx = seq_to_page[seq_i]
            top_candidates.append({
                "content_seq": seq_i,
                "source_pdf_page": page_idx + 1,
                "title": page_features[page_idx]["title"],
                "score": round(score, 4),
            })

        diagnostics.append({
            "block_index_1based": j + 1,
            "timecode": block.timecode,
            "hinted_content_seq": block.hinted_page,
            "assigned_content_seq": assigned_seq,
            "assigned_source_pdf_page": assigned_page_idx + 1,
            "assigned_slide_title": page_features[assigned_page_idx]["title"],
            "assigned_score": round(assigned_score, 4),
            "best_score": round(best_score, 4),
            "second_score": round(second_score, 4),
            "margin_vs_best": round(margin_vs_best, 4),
            "margin_vs_second": round(margin_vs_second, 4),
            "priority_hit_count": priority_hits,
            "title_hit_count": title_hits,
            "core_hit_count": core_hits,
            "confidence": confidence,
            "review_reasons": reasons,
            "top_candidates": top_candidates,
            "summary": block_features[j]["summary"],
            "medical_keywords": block_features[j].get("medical_keywords", []),
        })

    summary = {
        "total_blocks": len(diagnostics),
        "high_confidence_blocks": sum(1 for d in diagnostics if d["confidence"] == "high"),
        "medium_confidence_blocks": sum(1 for d in diagnostics if d["confidence"] == "medium"),
        "low_confidence_blocks": sum(1 for d in diagnostics if d["confidence"] == "low"),
        "needs_manual_review": any(d["confidence"] == "low" for d in diagnostics),
        "low_confidence_block_indices": [d["block_index_1based"] for d in diagnostics if d["confidence"] == "low"],
    }
    return diagnostics, summary


def add_underlines_to_source_page(page: fitz.Page, page_feat: Dict[str, object], block_feats: List[Dict[str, object]]) -> List[str]:
    overlap: List[str] = []
    page_kw = {kw.lower(): kw for kw in (page_feat.get("priority_keywords", []) + page_feat["title_keywords"] + page_feat["summary_keywords"] + page_feat["body_keywords"])}
    block_kw: Dict[str, str] = {}
    for bf in block_feats:
        for kw in bf["body_keywords"]:
            block_kw.setdefault(kw.lower(), kw)
    common = [page_kw[k] for k in page_kw if k in block_kw]
    common = sorted(set(common), key=lambda s: (-len(s), s.lower()))
    for kw in common:
        try:
            rects = page.search_for(kw)
        except Exception:
            rects = []
        for rect in rects[:40]:
            y = rect.y1 + 1.5
            page.draw_line((rect.x0, y), (rect.x1, y), color=(1.0, 0.55, 0.0), width=1.8)
        if rects:
            overlap.append(kw)
    return overlap


def load_pil_font(fontfile: Optional[str], size: int):
    try:
        if fontfile and os.path.exists(fontfile):
            return ImageFont.truetype(fontfile, size=size)
    except Exception:
        pass
    return ImageFont.load_default()


def measure(draw: ImageDraw.ImageDraw, text: str, font) -> Tuple[int, int]:
    if not text:
        return 0, 0
    bbox = draw.textbbox((0, 0), text, font=font)
    return bbox[2] - bbox[0], bbox[3] - bbox[1]


def wrap_text_for_pil(text: str, font, max_width: int) -> List[str]:
    img = Image.new("RGB", (max_width, 10), "white")
    draw = ImageDraw.Draw(img)
    result: List[str] = []
    paragraphs = text.split("\n")
    for para in paragraphs:
        para = para.strip()
        if not para:
            result.append("")
            continue
        words = para.split(" ")
        current = words[0]
        for word in words[1:]:
            candidate = current + " " + word
            if measure(draw, candidate, font)[0] <= max_width:
                current = candidate
            else:
                if measure(draw, word, font)[0] > max_width:
                    chunk = ""
                    for ch in word:
                        cand = chunk + ch
                        if measure(draw, cand, font)[0] <= max_width:
                            chunk = cand
                        else:
                            if chunk:
                                result.append(current)
                                current = chunk
                            chunk = ch
                    current = current + " " + chunk if current else chunk
                else:
                    result.append(current)
                    current = word
        result.append(current)
    return result


def build_transcript_summary_table(related_blocks: List[TranscriptBlock], related_feats: List[Dict[str, object]], page_feat: Optional[Dict[str, object]] = None) -> str:
    if not related_blocks:
        return "항목 | 정리\n상태 | 관련 전사문 없음"

    rows = build_compact_concept_rows(
        related_blocks,
        related_feats,
        page_feat or {},
        max_rows=4,
        max_value_chars=38,
        include_overview=True,
    )
    return rows_to_table_text(rows)


def paginate_lines(lines: List[str], font, max_height: int, max_width: int, line_spacing: int = 6) -> List[List[str]]:
    img = Image.new("RGB", (max_width, max_height), "white")
    draw = ImageDraw.Draw(img)
    _, line_h = measure(draw, "가A", font)
    line_h = max(line_h, 14) + line_spacing
    max_lines = max(1, max_height // line_h)
    chunks: List[List[str]] = []
    for i in range(0, len(lines), max_lines):
        chunks.append(lines[i:i + max_lines])
    return chunks or [[""]]


def choose_font_and_paginate(
    text: str,
    width_px: int,
    height_px: int,
    fontfile: Optional[str],
    preferred_font_size: int,
    min_font_size: int = 11,
    preferred_max_pages: int = 3,
) -> Tuple[List[List[str]], int]:
    for fs in range(preferred_font_size, min_font_size - 1, -1):
        font = load_pil_font(fontfile, fs)
        lines = wrap_text_for_pil(text, font, max_width=width_px - 24)
        chunks = paginate_lines(lines, font, max_height=height_px - 24, max_width=width_px - 24)
        if len(chunks) <= preferred_max_pages:
            return chunks, fs
    fs = min_font_size
    font = load_pil_font(fontfile, fs)
    lines = wrap_text_for_pil(text, font, max_width=width_px - 24)
    chunks = paginate_lines(lines, font, max_height=height_px - 24, max_width=width_px - 24)
    return chunks, fs


def render_lines_panel_to_png(lines: List[str], width_px: int, height_px: int, fontfile: Optional[str], font_size: int, out_path: str) -> None:
    bg = Image.new("RGB", (max(width_px, 50), max(height_px, 50)), "white")
    draw = ImageDraw.Draw(bg)
    draw.rectangle((0, 0, width_px - 1, height_px - 1), outline=(220, 220, 220), width=2)
    font = load_pil_font(fontfile, font_size)
    y = 12
    _, line_h = measure(draw, "가A", font)
    line_h = max(line_h, 14) + 6
    max_y = height_px - 12
    for line in lines:
        if y + line_h > max_y:
            break
        draw.text((12, y), line, font=font, fill=(0, 0, 0))
        y += line_h
    bg.save(out_path)


DIGEST_HEADER_COLORS = [
    (228, 239, 251),
    (234, 244, 232),
    (247, 238, 221),
    (238, 232, 247),
]
SPECIAL_MARKER_ALIASES = {"PY", "P Y", "P.Y", "P-Y"}
SPECIAL_MARKER_LABEL = "PY"
DIGEST_BLUE_RGB = (34, 89, 196)
DIGEST_BLUE_SUB_RGB = (61, 105, 189)
DIGEST_BLUE_BADGE_FILL = (234, 242, 255)
DIGEST_BLUE_BADGE_OUTLINE = (171, 196, 238)
DIGEST_BLUE_BADGE_TEXT = (42, 86, 164)
DIGEST_RED_RGB = (191, 54, 68)
DIGEST_RED_SUB_RGB = (168, 65, 78)
DIGEST_PURPLE_RGB = (124, 65, 176)
DIGEST_PURPLE_SUB_RGB = (108, 76, 166)
DIGEST_PURPLE_BADGE_FILL = (242, 234, 255)
DIGEST_PURPLE_BADGE_OUTLINE = (205, 182, 242)
DIGEST_PURPLE_BADGE_TEXT = (108, 76, 166)


def normalize_marker_alias(text: str) -> str:
    raw = re.sub(r"[^A-Za-z]", "", str(text or "")).upper()
    return raw


def has_special_marker_hint(page: fitz.Page, color_priority: Optional[Dict[str, object]] = None) -> bool:
    color_priority = color_priority or {}
    text_candidates: List[str] = []
    for key in ("green_text_spans", "green_highlight_texts", "priority_fragments"):
        for item in color_priority.get(key, []) or []:
            s = re.sub(r"\s+", " ", str(item)).strip()
            if s:
                text_candidates.append(s)
    for item in text_candidates:
        normalized = normalize_marker_alias(item)
        if normalized == SPECIAL_MARKER_LABEL:
            return True
        if re.search(r"\bP\s*Y\b", item, flags=re.IGNORECASE):
            return True

    try:
        annot = page.first_annot
        while annot:
            info = getattr(annot, "info", {}) or {}
            contents = " ".join([str(info.get("content", "") or ""), str(info.get("title", "") or ""), str(info.get("subject", "") or "")]).strip()
            if normalize_marker_alias(contents) == SPECIAL_MARKER_LABEL or re.search(r"\bP\s*Y\b", contents, flags=re.IGNORECASE):
                return True
            annot = annot.next
    except Exception:
        pass

    page_rect = page.rect
    search_limit = fitz.Rect(page_rect.x0, page_rect.y0, page_rect.x0 + page_rect.width * 0.62, page_rect.y0 + page_rect.height * 0.36)
    try:
        annot = page.first_annot
        while annot:
            subtype = ""
            try:
                subtype = (annot.type[1] or "").lower()
            except Exception:
                subtype = ""
            colors = getattr(annot, "colors", {}) or {}
            stroke = colors.get("stroke")
            rect = annot.rect
            if subtype in {"ink", "freetext"} and rect and rect.intersects(search_limit) and _matches_specific_green(stroke, mode="handwriting"):
                if rect.width >= page_rect.width * 0.16 and rect.height >= page_rect.height * 0.10:
                    return True
            annot = annot.next
    except Exception:
        pass

    try:
        for draw_item in page.get_drawings():
            rect = _rect_from_drawing(draw_item)
            if not rect or not rect.intersects(search_limit):
                continue
            stroke = draw_item.get("color") or draw_item.get("stroke")
            fill = draw_item.get("fill")
            if not (_matches_specific_green(stroke, mode="handwriting") or _matches_specific_green(fill, mode="handwriting")):
                continue
            if rect.width >= page_rect.width * 0.16 and rect.height >= page_rect.height * 0.10:
                return True
    except Exception:
        pass

    return False


def _normalize_for_match(text: str) -> str:
    text = re.sub(r"\s+", " ", str(text or "")).strip().lower()
    return re.sub(r"[^0-9a-z가-힣]+", "", text)


def _fragment_matches_text(fragment: str, text: str) -> bool:
    frag_n = _normalize_for_match(fragment)
    text_n = _normalize_for_match(text)
    if not frag_n or not text_n:
        return False
    if frag_n in text_n or text_n in frag_n:
        return True
    frag_toks = [t.lower() for t in tokenize(fragment) if len(t) >= 2]
    text_toks = {t.lower() for t in tokenize(text) if len(t) >= 2}
    if not frag_toks or not text_toks:
        return False
    overlap = len(set(frag_toks) & text_toks)
    return overlap >= max(1, min(2, len(set(frag_toks))))


def _row_style(base_style: str, is_special_marker_page: bool) -> str:
    base_style = (base_style or "default").lower()
    if is_special_marker_page and base_style == "red":
        return "purple"
    if is_special_marker_page and base_style == "default":
        return "blue"
    return base_style


def split_sentences_for_digest(text: str) -> List[str]:
    text = re.sub(r"\s+", " ", text).strip()
    if not text:
        return []
    parts = re.split(r"(?<=[.!?])\s+|(?<=다)\s+|(?<=됨)\s+|(?<=있음)\s+|(?<=함)\s+", text)
    cleaned: List[str] = []
    for part in parts:
        part = re.sub(r"^[•▪\-\d\)\.\(\s]+", "", part).strip()
        if len(part) < 8:
            continue
        cleaned.append(part)
    return _dedupe_keep_order(cleaned)


def compact_join(parts: List[str], max_chars: int, sep: str = " / ") -> str:
    cleaned = []
    for part in parts:
        part = re.sub(r"\s+", " ", part).strip()
        if not part:
            continue
        cleaned.append(part)
    cleaned = _dedupe_keep_order(cleaned)
    if not cleaned:
        return "없음"
    out = cleaned[0]
    for part in cleaned[1:]:
        candidate = out + sep + part
        if len(candidate) > max_chars:
            break
        out = candidate
    if len(out) > max_chars:
        out = out[: max_chars - 1].rstrip() + "…"
    return out


def safe_text(value: str, max_chars: int) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    if not value:
        return "없음"
    if len(value) <= max_chars:
        return value
    return value[: max_chars - 1].rstrip() + "…"


ROMAN_SECTION_RE = re.compile(r"^\s*([IVXLC]+)\.\s*(.+)$")
TOPIC_PREFIX_RE = re.compile(r"^\s*(?:[IVXLC]+|\d+(?:[-.]\d+)?)\.\s*")


def normalize_topic_title(title: str) -> str:
    title = re.sub(r"\s+", " ", str(title or "")).strip()
    title = TOPIC_PREFIX_RE.sub("", title).strip(" -:•·")
    return title or "핵심 정리"


def prettify_keyword_label(keyword: str) -> str:
    keyword = str(keyword or "").strip()
    if not keyword:
        return "핵심"
    keyword = keyword.replace("_", " ")
    if re.fullmatch(r"[A-Z0-9\-]{2,}", keyword):
        return keyword
    if re.search(r"[A-Za-z]", keyword):
        words = keyword.split()
        return " ".join(w.upper() if len(w) <= 4 else (w[0].upper() + w[1:]) for w in words)
    return keyword


def clean_spoken_summary_text(text: str) -> str:
    text = re.sub(r"\s+", " ", str(text or "")).strip()
    if not text:
        return ""
    text = re.sub(r"^(그리고|그다음|이제|여기서|보시면|보면|일단|사실은|그러면)\s+", "", text)
    text = re.sub(r"(같아요|이렇게 보면|라고 할 수 있어요|라고 보면 돼요)", "", text)
    text = re.sub(r"\s+", " ", text).strip(" ,;:/")
    return text


def build_overview_sentence(related_blocks: List[TranscriptBlock], related_feats: List[Dict[str, object]], page_feat: Dict[str, object], max_chars: int = 64) -> str:
    first_sentence_pool: List[str] = []
    detail_sentence_pool: List[str] = []
    for feat, block in zip(related_feats, related_blocks):
        first_sentence = clean_spoken_summary_text(str(feat.get("first_sentence", "")))
        if first_sentence:
            first_sentence_pool.append(first_sentence)
        for sent in split_sentences_for_digest(block.body)[:3]:
            cleaned = clean_spoken_summary_text(sent)
            if cleaned:
                detail_sentence_pool.append(cleaned)
    core_bullets = [clean_spoken_summary_text(str(x)) for x in page_feat.get("core_bullets", []) if str(x).strip()]
    green_fragments = [clean_spoken_summary_text(str(x)) for x in page_feat.get("priority_fragments", []) if str(x).strip()]
    candidates = _dedupe_keep_order(first_sentence_pool + detail_sentence_pool + core_bullets + green_fragments)
    candidates = [x for x in candidates if len(x) >= 8]
    if not candidates:
        return "핵심 내용 정리"
    return compact_join(candidates[:3], max_chars=max_chars, sep=" / ")


def _keyword_match_score(keyword: str, sentence: str) -> int:
    kw_toks = {t.lower() for t in tokenize(keyword) if len(t) >= 2}
    sent_toks = {t.lower() for t in tokenize(sentence) if len(t) >= 2}
    if not kw_toks or not sent_toks:
        return 0
    return len(kw_toks & sent_toks)


def describe_keyword_from_sentences(keyword: str, sentences: List[str], max_chars: int = 56) -> str:
    keyword_norm = normalize_medical_terms(keyword)
    best = ""
    best_score = 0
    for sentence in sentences:
        cleaned = clean_spoken_summary_text(sentence)
        if not cleaned:
            continue
        score = _keyword_match_score(keyword_norm, cleaned)
        if score <= 0:
            continue
        candidate = cleaned
        if re.search(r"[A-Za-z]", keyword_norm):
            candidate = re.sub(r"(?i)\b" + re.escape(keyword_norm.replace("_", " ")) + r"\b", "", candidate)
        else:
            candidate = candidate.replace(keyword, "")
        candidate = re.sub(r"^[,:;·•\-\s]+", "", candidate)
        candidate = re.sub(r"\s+", " ", candidate).strip()
        if ":" in candidate:
            after = candidate.split(":", 1)[1].strip()
            if len(after) >= 4:
                candidate = after
        parts = [p.strip() for p in re.split(r"[;,]|\s+-\s+|\s+→\s+", candidate) if p.strip()]
        if parts:
            candidate = compact_join(parts[:3], max_chars=max_chars, sep=", ")
        candidate = safe_text(candidate or cleaned, max_chars)
        if score > best_score or (score == best_score and (not best or len(candidate) < len(best))):
            best = candidate
            best_score = score
    return best


def build_compact_concept_rows(
    related_blocks: List[TranscriptBlock],
    related_feats: List[Dict[str, object]],
    page_feat: Dict[str, object],
    max_rows: int = 4,
    max_value_chars: int = 56,
    include_overview: bool = True,
) -> List[Dict[str, str]]:
    if not related_blocks:
        return [{"label": "상태", "value": "관련 전사문 없음", "style": "default"}]

    is_special_marker_page = bool(page_feat.get("special_marker_page"))
    red_fragments = _dedupe_keep_order([str(x) for x in page_feat.get("red_fragments", []) if str(x).strip()])[:4]
    red_keywords = {str(x).lower() for x in page_feat.get("red_keywords", [])}

    sentences: List[str] = []
    for block in related_blocks:
        sentences.extend(split_sentences_for_digest(block.body))
    sentences = _dedupe_keep_order([clean_spoken_summary_text(s) for s in sentences if clean_spoken_summary_text(s)])

    rows: List[Dict[str, str]] = []
    title_label = normalize_topic_title(str(page_feat.get("title") or page_feat.get("headline_text") or "핵심"))
    if include_overview:
        overview = build_overview_sentence(related_blocks, related_feats, page_feat, max_chars=max_value_chars)
        overview_has_red = any(_fragment_matches_text(frag, overview) for frag in red_fragments)
        rows.append({
            "label": safe_text(title_label, 18),
            "value": safe_text(overview, max_value_chars),
            "style": _row_style("red" if overview_has_red else "default", is_special_marker_page),
        })

    keyword_pool: List[str] = []
    for source in [page_feat.get("priority_keywords", []), page_feat.get("title_keywords", []), page_feat.get("summary_keywords", []), page_feat.get("page_term_hints", []), page_feat.get("body_keywords", [])]:
        keyword_pool.extend([str(x) for x in source])
    for feat in related_feats:
        keyword_pool.extend([str(x) for x in feat.get("summary_keywords", [])])
        keyword_pool.extend([str(x) for x in feat.get("body_keywords", [])[:4]])
    keyword_pool = _dedupe_keep_order(keyword_pool)

    used_labels = {rows[0]["label"].lower()} if rows else set()
    for keyword in keyword_pool:
        label = prettify_keyword_label(keyword)
        label_key = normalize_medical_terms(label)
        if not label_key or label_key in used_labels:
            continue
        if label_key in STOPWORDS or len(label) < 2:
            continue
        if label_key == normalize_medical_terms(title_label):
            continue
        desc = describe_keyword_from_sentences(keyword, sentences, max_chars=max_value_chars)
        if not desc or desc == "없음":
            continue
        base_style = "red" if (label_key in red_keywords or any(_fragment_matches_text(frag, label + " " + desc) for frag in red_fragments)) else "default"
        rows.append({
            "label": safe_text(label, 20),
            "value": safe_text(desc, max_value_chars),
            "style": _row_style(base_style, is_special_marker_page),
        })
        used_labels.add(label_key)
        if len(rows) >= max_rows:
            break

    if len(rows) < max_rows and red_fragments:
        joined_text = " ".join(row["value"] for row in rows)
        for frag in red_fragments:
            if _fragment_matches_text(frag, joined_text):
                continue
            rows.append({
                "label": "보강",
                "value": safe_text(clean_spoken_summary_text(frag), max_value_chars),
                "style": _row_style("red", is_special_marker_page),
            })
            if len(rows) >= max_rows:
                break

    return rows[:max_rows] or [{"label": "핵심", "value": "관련 전사문 없음", "style": _row_style("default", is_special_marker_page)}]


def rows_to_table_text(rows: List[Dict[str, str]]) -> str:
    if not rows:
        return "항목 | 정리\n상태 | 관련 전사문 없음"
    lines = ["항목 | 정리"]
    for row in rows:
        lines.append(f"{row.get('label','항목')} | {row.get('value','없음')}")
    return "\n".join(lines)

def detect_major_section_title(title: str) -> Optional[str]:
    title = re.sub(r"\s+", " ", str(title or "")).strip()
    if not title:
        return None
    m = ROMAN_SECTION_RE.match(title)
    if m:
        return f"{m.group(1)}. {m.group(2).strip()}"
    return None


def group_entries_by_section(entries: List[Dict[str, object]], document_title: str) -> List[Tuple[str, List[Dict[str, object]]]]:
    groups: List[Tuple[str, List[Dict[str, object]]]] = []
    current_section = document_title
    bucket: List[Dict[str, object]] = []
    for entry in entries:
        section_hint = detect_major_section_title(str(entry.get("raw_title", "") or entry.get("title", "")))
        if section_hint and bucket:
            groups.append((current_section, bucket))
            bucket = []
            current_section = section_hint
        elif section_hint:
            current_section = section_hint
        bucket.append(entry)
    if bucket:
        groups.append((current_section, bucket))
    return groups or [(document_title, entries)]


def build_digest_entry(src_index: int, related_blocks: List[TranscriptBlock], related_feats: List[Dict[str, object]], page_feat: Dict[str, object]) -> Dict[str, object]:
    title = safe_text(normalize_topic_title(str(page_feat.get("title") or page_feat.get("headline_text") or f"원본 슬라이드 {src_index + 1}")), 64)
    subtitle = safe_text(str(page_feat.get("subtitle") or page_feat.get("headline_text") or ""), 84)
    if subtitle == title or subtitle == "없음":
        subtitle = ""

    rows = build_compact_concept_rows(
        related_blocks,
        related_feats,
        page_feat,
        max_rows=5,
        max_value_chars=70,
        include_overview=True,
    )

    red_fragments = _dedupe_keep_order([str(x) for x in page_feat.get("red_fragments", []) if str(x).strip()])[:4]
    is_special_marker_page = bool(page_feat.get("special_marker_page"))
    return {
        "source_page": src_index + 1,
        "title": title,
        "raw_title": str(page_feat.get("title") or page_feat.get("headline_text") or title),
        "subtitle": subtitle,
        "rows": rows,
        "special_marker_page": is_special_marker_page,
        "marker_label": SPECIAL_MARKER_LABEL if is_special_marker_page else "",
        "has_red_emphasis": bool(red_fragments),
        "red_fragments": red_fragments,
        "red_missing_fragments": [],
    }


def build_transcript_digest_entries(
    blocks: List[TranscriptBlock],
    mapping: Dict[int, List[int]],
    block_features: List[Dict[str, object]],
    page_features: Dict[int, Dict[str, object]],
) -> List[Dict[str, object]]:
    entries: List[Dict[str, object]] = []
    for src_index in range(len(page_features)):
        related_idxs = mapping.get(src_index, [])
        if not related_idxs:
            continue
        related_blocks = [blocks[i] for i in related_idxs]
        related_feats = [block_features[i] for i in related_idxs]
        entries.append(build_digest_entry(src_index, related_blocks, related_feats, page_features[src_index]))

    if entries:
        return entries

    if not blocks:
        return [{
            "source_page": 0,
            "title": "전사문을 찾지 못함",
            "subtitle": "입력된 전사문 형식을 다시 확인하세요.",
            "rows": [{"label": "상태", "value": "관련 전사문 없음", "style": "default"}],
            "special_marker_page": False,
            "marker_label": "",
        }]

    fallback_feats = [build_block_features(b) for b in blocks]
    chunk_size = 3
    for offset in range(0, len(blocks), chunk_size):
        related_blocks = blocks[offset: offset + chunk_size]
        related_feats = fallback_feats[offset: offset + chunk_size]
        page_feat = {
            "title": related_blocks[0].source_name or f"전사문 묶음 {offset // chunk_size + 1}",
            "subtitle": "슬라이드 매칭 전 요약",
            "headline_text": related_blocks[0].source_name or "",
            "priority_fragments": [],
            "priority_keywords": [],
            "summary_keywords": [],
            "title_keywords": extract_keywords(related_blocks[0].source_name or "", limit=5),
            "core_bullets": [],
        }
        entry = build_digest_entry(offset // chunk_size, related_blocks, related_feats, page_feat)
        entry["special_marker_page"] = False
        entry["marker_label"] = ""
        entries.append(entry)
    return entries


def render_transcript_digest_pdf_legacy_image(
    entries: List[Dict[str, object]],
    output_path: str,
    fontfile: Optional[str],
    document_title: str,
    subtitle: str,
    output_dir: str,
) -> str:
    os.makedirs(output_dir, exist_ok=True)
    page_w_pt, page_h_pt = 595, 842
    img_w, img_h = 1190, 1684
    margin_x, margin_y = 58, 52
    header_h = 118
    footer_h = 34
    block_gap = 22
    blocks_per_page = 3

    title_font = load_pil_font(fontfile, 38)
    subtitle_font = load_pil_font(fontfile, 20)
    section_font = load_pil_font(fontfile, 26)
    block_title_font = load_pil_font(fontfile, 24)
    block_subtitle_font = load_pil_font(fontfile, 16)
    row_label_font = load_pil_font(fontfile, 16)
    row_value_font = load_pil_font(fontfile, 17)
    badge_font = load_pil_font(fontfile, 15)
    footer_font = load_pil_font(fontfile, 15)

    temp_dir = os.path.join(output_dir, "_digest_imgs")
    os.makedirs(temp_dir, exist_ok=True)

    pdf = fitz.open()
    marked_entries = [entry for entry in entries if entry.get("special_marker_page")]
    grouped_sections = group_entries_by_section(entries, document_title)
    sections: List[Tuple[str, str, List[Dict[str, object]]]] = [
        (section_title, subtitle, section_entries) for section_title, section_entries in grouped_sections
    ]
    if marked_entries:
        sections.append(
            (
                f"{SPECIAL_MARKER_LABEL} 표시 페이지 모아보기",
                f"{SPECIAL_MARKER_LABEL} 표시가 있는 페이지의 전사문 요약만 마지막에 한 번 더 모아 정리",
                marked_entries,
            )
        )

    total_pages = sum(max(1, (len(section_entries) + blocks_per_page - 1) // blocks_per_page) for _, _, section_entries in sections)
    absolute_page_no = 0

    def draw_block(draw: ImageDraw.ImageDraw, entry: Dict[str, object], x0: int, y0: int, x1: int, y1: int, color_seed: int) -> None:
        draw.rounded_rectangle((x0, y0, x1, y1), radius=20, fill=(255, 255, 255), outline=(214, 220, 227), width=2)
        header_color = DIGEST_HEADER_COLORS[color_seed % len(DIGEST_HEADER_COLORS)]
        header_h_block = 56
        draw.rounded_rectangle((x0, y0, x1, y0 + header_h_block), radius=20, fill=header_color, outline=(214, 220, 227), width=0)
        draw.rectangle((x0, y0 + header_h_block - 16, x1, y0 + header_h_block), fill=header_color)

        marker_label = str(entry.get("marker_label", "") or "").strip()
        is_special_marker_page = bool(entry.get("special_marker_page"))
        has_red_emphasis = bool(entry.get("has_red_emphasis"))
        badge_text = f"{marker_label} · 슬라이드 {entry['source_page']}" if (is_special_marker_page and entry.get('source_page')) else f"슬라이드 {entry.get('source_page', 0)}"
        badge_w, badge_h = measure(draw, badge_text, badge_font)
        bx1 = x1 - 16
        bx0 = bx1 - badge_w - 18
        by0 = y0 + 12
        by1 = by0 + badge_h + 10
        if is_special_marker_page and has_red_emphasis:
            badge_fill = DIGEST_PURPLE_BADGE_FILL
            badge_outline = DIGEST_PURPLE_BADGE_OUTLINE
            badge_text_fill = DIGEST_PURPLE_BADGE_TEXT
        else:
            badge_fill = DIGEST_BLUE_BADGE_FILL if is_special_marker_page else (255, 255, 255)
            badge_outline = DIGEST_BLUE_BADGE_OUTLINE if is_special_marker_page else (207, 214, 221)
            badge_text_fill = DIGEST_BLUE_BADGE_TEXT if is_special_marker_page else (88, 97, 108)
        draw.rounded_rectangle((bx0, by0, bx1, by1), radius=14, fill=badge_fill, outline=badge_outline, width=1)
        draw.text((bx0 + 9, by0 + 5), badge_text, font=badge_font, fill=badge_text_fill)

        if is_special_marker_page and has_red_emphasis:
            title_fill = DIGEST_PURPLE_RGB
            subtitle_fill = DIGEST_PURPLE_SUB_RGB
        elif is_special_marker_page:
            title_fill = DIGEST_BLUE_RGB
            subtitle_fill = DIGEST_BLUE_SUB_RGB
        elif has_red_emphasis:
            title_fill = DIGEST_RED_RGB
            subtitle_fill = DIGEST_RED_SUB_RGB
        else:
            title_fill = (43, 56, 72)
            subtitle_fill = (102, 114, 128)

        title_max_width = x1 - x0 - 38 - (bx1 - bx0) - 12
        title_lines = wrap_text_for_pil(str(entry.get("title", "")), block_title_font, max_width=title_max_width)[:2]
        if len(title_lines) == 2 and len(title_lines[-1]) > 1:
            title_lines[-1] = safe_text(title_lines[-1], 26)
        ty = y0 + 12
        _, title_lh = measure(draw, "가A", block_title_font)
        title_lh = max(title_lh, 18) + 3
        for line in title_lines:
            draw.text((x0 + 16, ty), line, font=block_title_font, fill=title_fill)
            ty += title_lh

        subtitle_text = str(entry.get("subtitle", "")).strip()
        body_top = y0 + header_h_block + 12
        if subtitle_text:
            subtitle_text = safe_text(subtitle_text, 86)
            draw.text((x0 + 16, body_top), subtitle_text, font=block_subtitle_font, fill=subtitle_fill)
            body_top += 24

        table_x0 = x0 + 14
        table_x1 = x1 - 14
        label_col_w = 126
        current_y = body_top
        max_bottom = y1 - 16
        for idx, row in enumerate(entry.get("rows", [])):
            label = safe_text(str(row.get("label", "항목")), 24)
            value = safe_text(str(row.get("value", "없음")), 110)
            row_style = str(row.get("style", "default") or "default").lower()
            label_lines = wrap_text_for_pil(label, row_label_font, max_width=label_col_w - 18)
            value_lines = wrap_text_for_pil(value, row_value_font, max_width=(table_x1 - table_x0 - label_col_w - 24))[:3]
            label_h = max(1, len(label_lines)) * (max(measure(draw, "가A", row_label_font)[1], 14) + 4)
            value_h = max(1, len(value_lines)) * (max(measure(draw, "가A", row_value_font)[1], 14) + 4)
            row_h = max(label_h, value_h) + 14
            if current_y + row_h > max_bottom:
                break
            row_fill = (248, 250, 252) if idx % 2 == 0 else (255, 255, 255)
            draw.rounded_rectangle((table_x0, current_y, table_x1, current_y + row_h), radius=10, fill=row_fill, outline=(227, 231, 236), width=1)
            draw.line((table_x0 + label_col_w, current_y + 6, table_x0 + label_col_w, current_y + row_h - 6), fill=(224, 228, 233), width=1)
            ly = current_y + 8
            for line in label_lines:
                draw.text((table_x0 + 10, ly), line, font=row_label_font, fill=(68, 79, 94))
                ly += max(measure(draw, "가A", row_label_font)[1], 14) + 4
            if row_style == "purple":
                value_fill = DIGEST_PURPLE_RGB
            elif row_style == "red":
                value_fill = DIGEST_RED_RGB
            elif row_style == "blue":
                value_fill = DIGEST_BLUE_RGB
            else:
                value_fill = (34, 34, 34)
            vy = current_y + 8
            for line in value_lines:
                draw.text((table_x0 + label_col_w + 10, vy), line, font=row_value_font, fill=value_fill)
                vy += max(measure(draw, "가A", row_value_font)[1], 14) + 4
            current_y += row_h + 8

    for section_title, section_subtitle, section_entries in sections:
        section_total = max(1, (len(section_entries) + blocks_per_page - 1) // blocks_per_page)
        for section_page_idx in range(section_total):
            bg = Image.new("RGB", (img_w, img_h), (248, 249, 251))
            draw = ImageDraw.Draw(bg)
            title_y = margin_y
            draw.text((margin_x, title_y), document_title, font=title_font, fill=(34, 52, 74))
            draw.text((margin_x, title_y + 46), section_title, font=section_font, fill=(58, 73, 92))
            draw.text((margin_x, title_y + 82), section_subtitle, font=subtitle_font, fill=(94, 107, 125))
            page_label = f"{absolute_page_no + 1} / {total_pages}"
            absolute_page_no += 1
            label_w, label_h = measure(draw, page_label, subtitle_font)
            draw.rounded_rectangle((img_w - margin_x - label_w - 24, title_y + 10, img_w - margin_x, title_y + 10 + label_h + 14), radius=16, fill=(232, 237, 243), outline=(210, 217, 225), width=1)
            draw.text((img_w - margin_x - label_w - 12, title_y + 17), page_label, font=subtitle_font, fill=(77, 89, 102))

            start = section_page_idx * blocks_per_page
            page_entries = section_entries[start:start + blocks_per_page]
            usable_top = margin_y + header_h
            usable_bottom = img_h - footer_h - 18
            block_h = int((usable_bottom - usable_top - block_gap * (blocks_per_page - 1)) / blocks_per_page)
            for local_idx, entry in enumerate(page_entries):
                y0 = usable_top + local_idx * (block_h + block_gap)
                y1 = y0 + block_h
                draw_block(draw, entry, margin_x, y0, img_w - margin_x, y1, start + local_idx)

            footer = "자동 생성 요약본 - 강의록 구조를 따라 전사문을 표형 정리로 재구성"
            if section_title.startswith(SPECIAL_MARKER_LABEL):
                footer = f"자동 생성 요약본 - {SPECIAL_MARKER_LABEL} 표시 페이지 모아보기"
            draw.text((margin_x, img_h - footer_h), footer, font=footer_font, fill=(116, 125, 135))

            img_path = os.path.join(temp_dir, f"digest_page_{absolute_page_no}.png")
            bg.save(img_path)
            page = pdf.new_page(width=page_w_pt, height=page_h_pt)
            page.insert_image(fitz.Rect(0, 0, page_w_pt, page_h_pt), filename=img_path)

    pdf.save(output_path)
    pdf.close()
    return output_path



def _set_docx_cell_background(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill_hex)


def _set_docx_cell_width(cell, width_inch: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    twips = int(width_inch * 1440)
    tc_w.set(qn('w:w'), str(twips))
    tc_w.set(qn('w:type'), 'dxa')


def _append_digest_value_run(paragraph, text: str, style_name: str = 'default'):
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    if style_name == 'blue':
        run.font.color.rgb = RGBColor(*DIGEST_BLUE_RGB)
    elif style_name == 'red':
        run.font.color.rgb = RGBColor(*DIGEST_RED_RGB)
    elif style_name == 'purple':
        run.font.color.rgb = RGBColor(*DIGEST_PURPLE_RGB)
    else:
        run.font.color.rgb = RGBColor(34, 34, 34)
    return run


def render_transcript_digest_docx(
    entries: List[Dict[str, object]],
    output_path: str,
    document_title: str,
    subtitle: str,
) -> str:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Noto Sans CJK KR'
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Noto Sans CJK KR')
    normal_style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    r = p.add_run(document_title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(34, 52, 74)

    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(94, 107, 125)

    grouped_sections = group_entries_by_section(entries, document_title)
    marked_entries = [entry for entry in entries if entry.get('special_marker_page')]
    sections: List[Tuple[str, str, List[Dict[str, object]]]] = [
        (section_title, subtitle, section_entries) for section_title, section_entries in grouped_sections
    ]
    if marked_entries:
        sections.append((
            f'{SPECIAL_MARKER_LABEL} 표시 페이지 모아보기',
            f'{SPECIAL_MARKER_LABEL} 표시가 있는 페이지의 전사문 요약만 마지막에 한 번 더 모아 정리',
            marked_entries,
        ))

    for sec_idx, (section_title, section_subtitle, section_entries) in enumerate(sections):
        if sec_idx > 0:
            doc.add_section(WD_SECTION_START.NEW_PAGE)
        hp = doc.add_paragraph()
        rr = hp.add_run(section_title)
        rr.bold = True
        rr.font.size = Pt(14)
        rr.font.color.rgb = RGBColor(58, 73, 92)

        sp = doc.add_paragraph()
        sr = sp.add_run(section_subtitle)
        sr.font.size = Pt(9.5)
        sr.font.color.rgb = RGBColor(94, 107, 125)

        for entry in section_entries:
            badge_text = f"슬라이드 {entry.get('source_page', 0)}"
            if entry.get('special_marker_page') and entry.get('marker_label'):
                badge_text = f"{entry.get('marker_label')} · {badge_text}"

            title_p = doc.add_paragraph()
            title_r = title_p.add_run(f"[{badge_text}] {entry.get('title', '')}")
            title_r.bold = True
            title_r.font.size = Pt(11.5)
            if entry.get('special_marker_page') and entry.get('has_red_emphasis'):
                title_r.font.color.rgb = RGBColor(*DIGEST_PURPLE_RGB)
            elif entry.get('special_marker_page'):
                title_r.font.color.rgb = RGBColor(*DIGEST_BLUE_RGB)
            elif entry.get('has_red_emphasis'):
                title_r.font.color.rgb = RGBColor(*DIGEST_RED_RGB)
            else:
                title_r.font.color.rgb = RGBColor(43, 56, 72)

            subtitle_text = str(entry.get('subtitle', '') or '').strip()
            if subtitle_text:
                sub_p = doc.add_paragraph()
                sub_p.paragraph_format.space_after = Pt(4)
                sub_r = sub_p.add_run(subtitle_text)
                sub_r.italic = True
                sub_r.font.size = Pt(9.5)
                if entry.get('special_marker_page') and entry.get('has_red_emphasis'):
                    sub_r.font.color.rgb = RGBColor(*DIGEST_PURPLE_SUB_RGB)
                elif entry.get('special_marker_page'):
                    sub_r.font.color.rgb = RGBColor(*DIGEST_BLUE_SUB_RGB)
                elif entry.get('has_red_emphasis'):
                    sub_r.font.color.rgb = RGBColor(*DIGEST_RED_SUB_RGB)
                else:
                    sub_r.font.color.rgb = RGBColor(102, 114, 128)

            rows = entry.get('rows', []) or [{'label': '상태', 'value': '관련 전사문 없음', 'style': 'default'}]
            table = doc.add_table(rows=1, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            hdr[0].text = '항목'
            hdr[1].text = '정리'
            for c in hdr:
                _set_docx_cell_background(c, 'EAEFF6')
                for p0 in c.paragraphs:
                    for run in p0.runs:
                        run.bold = True
                        run.font.size = Pt(10.5)
                        run.font.color.rgb = RGBColor(48, 63, 82)
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_docx_cell_width(hdr[0], 1.45)
            _set_docx_cell_width(hdr[1], 5.9)

            for row in rows:
                label = safe_text(str(row.get('label', '항목')), 40)
                value = safe_text(str(row.get('value', '없음')), 500)
                style_name = str(row.get('style', 'default') or 'default').lower()
                cells = table.add_row().cells
                cells[0].text = label
                cells[1].text = ''
                _set_docx_cell_width(cells[0], 1.45)
                _set_docx_cell_width(cells[1], 5.9)
                cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                _set_docx_cell_background(cells[0], 'F6F8FB')
                for p0 in cells[0].paragraphs:
                    for run in p0.runs:
                        run.bold = True
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(68, 79, 94)
                vp = cells[1].paragraphs[0]
                _append_digest_value_run(vp, value, style_name=style_name)

            doc.add_paragraph()

    doc.save(output_path)
    return output_path


def convert_docx_to_pdf_with_libreoffice(docx_path: str, output_dir: str) -> Tuple[Optional[str], Dict[str, object]]:
    exe = shutil.which('soffice') or shutil.which('libreoffice')
    info: Dict[str, object] = {
        'docx_path': docx_path,
        'output_dir': output_dir,
        'converter': exe,
        'success': False,
        'stdout': '',
        'stderr': '',
        'verified_page_count': 0,
    }
    if not exe:
        info['stderr'] = 'LibreOffice executable not found'
        return None, info

    outdir = Path(output_dir)
    outdir.mkdir(parents=True, exist_ok=True)
    expected_pdf = outdir / (Path(docx_path).stem + '.pdf')
    temp_home = tempfile.mkdtemp(prefix='lo_headless_home_')
    profile_dir = Path(temp_home) / 'lo_profile'
    profile_dir.mkdir(parents=True, exist_ok=True)
    try:
        cmd = [
            exe,
            '--headless',
            '--nologo',
            '--nodefault',
            '--nolockcheck',
            '--norestore',
            f'-env:UserInstallation=file://{profile_dir.as_posix()}',
            '--convert-to', 'pdf:writer_pdf_Export',
            '--outdir', str(outdir),
            str(docx_path),
        ]
        env = os.environ.copy()
        env['HOME'] = temp_home
        env['SAL_USE_VCLPLUGIN'] = 'svp'
        env['DISPLAY'] = ''
        proc = subprocess.run(cmd, capture_output=True, text=True, env=env)
        info['stdout'] = proc.stdout
        info['stderr'] = proc.stderr
        info['returncode'] = proc.returncode
        if expected_pdf.exists() and expected_pdf.stat().st_size > 0:
            try:
                pdf = fitz.open(str(expected_pdf))
                info['verified_page_count'] = len(pdf)
                pdf.close()
                info['success'] = info['verified_page_count'] > 0
            except Exception as exc:
                info['stderr'] = (info.get('stderr') or '') + f'\nPDF verify failed: {exc}'
                info['success'] = False
        return (str(expected_pdf) if info['success'] else None), info
    finally:
        shutil.rmtree(temp_home, ignore_errors=True)


def render_transcript_digest_outputs(
    entries: List[Dict[str, object]],
    pdf_output_path: str,
    fontfile: Optional[str],
    document_title: str,
    subtitle: str,
    output_dir: str,
) -> Dict[str, object]:
    os.makedirs(output_dir, exist_ok=True)
    docx_output_path = os.path.join(output_dir, 'overall_transcript_summary_digest.docx')
    info: Dict[str, object] = {
        'docx_path': None,
        'pdf_path': None,
        'generation_mode': 'docx_first_then_pdf',
        'conversion': {},
        'fallback_used': False,
    }
    render_transcript_digest_docx(entries, docx_output_path, document_title, subtitle)
    info['docx_path'] = docx_output_path
    converted_pdf, conversion_info = convert_docx_to_pdf_with_libreoffice(docx_output_path, output_dir)
    info['conversion'] = conversion_info
    if converted_pdf and Path(converted_pdf).exists():
        final_pdf = Path(pdf_output_path)
        if Path(converted_pdf).resolve() != final_pdf.resolve():
            shutil.copy2(converted_pdf, final_pdf)
        info['pdf_path'] = str(final_pdf)
        return info

    render_transcript_digest_pdf_legacy_image(entries, pdf_output_path, fontfile, document_title, subtitle, output_dir)
    info['pdf_path'] = pdf_output_path
    info['fallback_used'] = True
    info['generation_mode'] = 'docx_first_then_pdf_with_image_pdf_fallback'
    return info

def build_rendering_coverage_report(
    blocks: List[TranscriptBlock],
    rendered_transcript_fragments: List[str],
    block_diagnostics: List[Dict[str, object]],
) -> Tuple[Dict[str, object], str]:
    original_text = "\n\n".join(f"[{b.timecode}] {b.body}" for b in blocks).strip()
    rendered_text = "\n\n".join(fragment for fragment in rendered_transcript_fragments if fragment.strip()).strip()

    norm_original = normalize_for_exact_compare(original_text)
    norm_rendered = normalize_for_exact_compare(rendered_text)

    low_conf_blocks = [d for d in block_diagnostics if d["confidence"] == "low"]
    medium_conf_blocks = [d for d in block_diagnostics if d["confidence"] == "medium"]

    coverage = {
        "original_transcript_char_count": len(original_text),
        "rendered_transcript_char_count": len(rendered_text),
        "normalized_original_char_count": len(norm_original),
        "normalized_rendered_char_count": len(norm_rendered),
        "normalized_text_exact_match": norm_original == norm_rendered,
        "original_transcript_sha1_normalized": sha1_text(norm_original),
        "rendered_transcript_sha1_normalized": sha1_text(norm_rendered),
        "all_blocks_rendered_once_in_order": norm_original == norm_rendered,
        "low_confidence_blocks": [d["block_index_1based"] for d in low_conf_blocks],
        "medium_confidence_blocks": [d["block_index_1based"] for d in medium_conf_blocks],
        "needs_manual_review": (norm_original != norm_rendered) or bool(low_conf_blocks),
        "manual_review_reasons": [reason for d in low_conf_blocks for reason in d.get("review_reasons", [])],
    }

    summary_lines = [
        f"원본 전사문 문자수: {coverage['original_transcript_char_count']}",
        f"렌더링된 전사문 문자수: {coverage['rendered_transcript_char_count']}",
        f"정규화 텍스트 완전일치: {coverage['normalized_text_exact_match']}",
        f"낮은 확신도 블록: {coverage['low_confidence_blocks']}",
        f"수동 검토 필요: {coverage['needs_manual_review']}",
    ]
    if coverage["manual_review_reasons"]:
        summary_lines.append("검토 이유: " + " / ".join(sorted(set(coverage["manual_review_reasons"]))))
    return coverage, "\n".join(summary_lines)


def _open_source_docs(source_pdf_map: Dict[str, str]) -> Dict[str, fitz.Document]:
    return {name: fitz.open(path) for name, path in source_pdf_map.items()}


def _build_single_pdf_plan(pdf_path: str) -> Tuple[List[LogicalPageRef], Dict[str, str]]:
    src = fitz.open(pdf_path)
    try:
        plan = [LogicalPageRef(source="single", page_no=i + 1, reason="single_original") for i in range(len(src))]
    finally:
        src.close()
    return plan, {"single": pdf_path}


def _generate_annotated_pdf_core(
    page_plan: List[LogicalPageRef],
    source_pdf_map: Dict[str, str],
    transcript_text: str,
    output_dir: str,
    render_mode_label: str,
) -> Tuple[str, str, str, str, str]:
    os.makedirs(output_dir, exist_ok=True)
    temp_img_dir = os.path.join(output_dir, "_panel_imgs")
    os.makedirs(temp_img_dir, exist_ok=True)

    precheck_path = os.path.join(output_dir, "verification_report_precheck.json")
    finalcheck_path = os.path.join(output_dir, "verification_report.json")
    completeness_json_path = os.path.join(output_dir, "transcript_completeness_review.json")
    completeness_txt_path = os.path.join(output_dir, "transcript_completeness_review.txt")
    output_pdf_path = os.path.join(output_dir, "annotated_output.pdf")
    digest_pdf_path = os.path.join(output_dir, "overall_transcript_summary_digest.pdf")

    fontfile = find_korean_font()
    blocks = parse_transcript(transcript_text)

    docs = _open_source_docs(source_pdf_map)
    try:
        logical_pages = []
        page_texts = []
        for ref in page_plan:
            page = docs[ref.source].load_page(ref.page_no - 1)
            logical_pages.append((ref, page))
            page_texts.append(page.get_text("text"))

        page_classes = [classify_page(txt, i) for i, txt in enumerate(page_texts)]
        seq_to_page, page_to_seq, content_indices = content_page_sequence_map(page_classes)
        transcript_anchor_seq_to_page, transcript_anchor_page_to_seq, transcript_anchor_indices = transcript_anchor_sequence_map(page_plan, page_classes)
        page_color_priorities = [extract_page_color_priority(logical_pages[i][1]) for i in range(len(page_texts))]
        page_features = {
            i: build_page_features(page_texts[i], page_color_priorities[i])
            for i in range(len(page_texts))
        }
        special_marker_pages = [has_special_marker_hint(logical_pages[i][1], page_color_priorities[i]) for i in range(len(page_texts))]
        for i in range(len(page_features)):
            page_features[i]["special_marker_page"] = special_marker_pages[i]

        mapping, metrics = align_blocks_to_pages(blocks, content_indices, page_features, page_to_seq, transcript_anchor_seq_to_page)
        block_features = [build_block_features(b) for b in blocks]
        digest_entries = build_transcript_digest_entries(blocks, mapping, block_features, page_features)
        block_diagnostics, diagnostic_summary = build_assignment_diagnostics(
            blocks,
            block_features,
            content_indices,
            page_features,
            page_to_seq,
            seq_to_page,
            transcript_anchor_seq_to_page,
            metrics.get("assignment_content_seq", []),
        )

        precheck = {
            "source_pdf_map": source_pdf_map,
            "logical_total_pages": len(logical_pages),
            "transcript_blocks": len(blocks),
            "content_pages": len(content_indices),
            "font_found": bool(fontfile),
            "fontfile": fontfile,
            "mapping_metrics": metrics,
            "assignment_diagnostic_summary": diagnostic_summary,
            "review_needed": (not bool(fontfile)) or (not bool(blocks)) or diagnostic_summary["needs_manual_review"],
            "render_source_policy": render_mode_label,
            "continuation_policy": "전사문은 길이에 상관없이 모든 내용을 유지하며 필요한 만큼 continuation page 생성",
            "matching_rule": "전사문 부착 매칭은 1순위 지정 초록색(#007355 필기, #A5C891 형광펜), 2순위 전사본 페이지 번호(줄 기준으로 해석하되 줄+야붙 병합 시 앞에 추가된 야붙 1페이지와 문제만 있는 페이지는 제외), 3순위 제목·부제·상단 내용·본문 훑기 결과를 함께 사용하며 한국어 설명은 가능한 범위에서 영문 의학용어 canonical form으로 정규화해 비교",
            "generated_transcript_digest_pdf": os.path.basename(digest_pdf_path),
            "generated_transcript_digest_docx": "overall_transcript_summary_digest.docx",
            "digest_entry_count": len(digest_entries),
            "digest_layout": "sectioned_portrait_summary_tables",
            "special_marker_label": SPECIAL_MARKER_LABEL,
            "special_marker_page_count": sum(1 for x in special_marker_pages if x),
            "special_marker_source_pages": [i + 1 for i, x in enumerate(special_marker_pages) if x],
            "red_emphasis_pages": [i + 1 for i, pf in page_features.items() if pf.get("red_fragments")],
        }
        with open(precheck_path, "w", encoding="utf-8") as f:
            json.dump(precheck, f, ensure_ascii=False, indent=2)

        out = fitz.open()
        all_overlaps: Dict[str, List[str]] = {}
        rendered_transcript_fragments: List[str] = []
        page_attachment_summary: List[Dict[str, object]] = []
        total_continuation_pages = 0

        transcript_font_size = 14
        note_font_size = 12

        for logical_index, (ref, src_page) in enumerate(logical_pages):
            related_block_idxs = mapping.get(logical_index, [])
            related_blocks = [blocks[k] for k in related_block_idxs]
            related_feats = [block_features[k] for k in related_block_idxs]

            transcript_raw = "\n\n".join(f"[{b.timecode}] {b.body}" for b in related_blocks).strip()
            transcript_summary_table = build_transcript_summary_table(related_blocks, related_feats, page_features[logical_index])
            top_title = page_features[logical_index]["title"]

            overlap_terms: List[str] = []
            if related_blocks and page_classes[logical_index] == "content":
                overlap_terms = add_underlines_to_source_page(src_page, page_features[logical_index], related_feats)
            all_overlaps[str(logical_index + 1)] = overlap_terms

            page_attachment_summary.append({
                "logical_output_page": logical_index + 1,
                "source_kind": ref.source,
                "source_pdf_page": ref.page_no,
                "page_class": page_classes[logical_index],
                "slide_title": top_title,
                "attached_block_indices": [idx + 1 for idx in related_block_idxs],
                "attached_timecodes": [blocks[idx].timecode for idx in related_block_idxs],
            })

            pix = src_page.get_pixmap(matrix=fitz.Matrix(1.8, 1.8), alpha=False, annots=True)
            page_width = src_page.rect.width
            page_height = src_page.rect.height
            extra_bottom = 290
            new_page = out.new_page(width=page_width, height=page_height + extra_bottom)
            new_page.insert_image(fitz.Rect(0, 0, page_width, page_height), pixmap=pix)

            y0 = page_height + 16
            x_margin = 22
            transcript_box = fitz.Rect(x_margin, y0, page_width * 0.68, page_height + extra_bottom - 18)
            note_box = fitz.Rect(page_width * 0.70, y0, page_width - 16, page_height + extra_bottom - 18)
            new_page.draw_rect(transcript_box, color=(0.85, 0.85, 0.85), width=0.8)
            new_page.draw_rect(note_box, color=(0.85, 0.85, 0.85), width=0.8)
            new_page.insert_text(
                (x_margin, y0 - 5),
                f"원본 슬라이드 {logical_index + 1}  [출처: {ref.source} {ref.page_no}]",
                fontsize=9.5,
                color=(0.22, 0.22, 0.22),
            )

            # Keep exact transcript order for completeness review.
            if transcript_raw:
                rendered_transcript_fragments.append(transcript_raw)

            transcript_display = transcript_raw if transcript_raw else "관련 전사문 없음"
            transcript_chunks, used_tr_font = choose_font_and_paginate(
                transcript_display,
                int(transcript_box.width),
                int(transcript_box.height),
                fontfile,
                transcript_font_size,
                min_font_size=11,
                preferred_max_pages=3,
            )

            tr_img_path = os.path.join(temp_img_dir, f"transcript_{logical_index+1}_1.png")
            render_lines_panel_to_png(
                transcript_chunks[0],
                int(transcript_box.width),
                int(transcript_box.height),
                fontfile,
                used_tr_font,
                tr_img_path,
            )
            new_page.insert_image(transcript_box, filename=tr_img_path)

            note_text = transcript_summary_table
            note_chunks, used_note_font = choose_font_and_paginate(
                note_text,
                int(note_box.width),
                int(note_box.height),
                fontfile,
                note_font_size,
                min_font_size=10,
                preferred_max_pages=1,
            )
            note_img_path = os.path.join(temp_img_dir, f"note_{logical_index+1}.png")
            render_lines_panel_to_png(
                note_chunks[0],
                int(note_box.width),
                int(note_box.height),
                fontfile,
                used_note_font,
                note_img_path,
            )
            new_page.insert_image(note_box, filename=note_img_path)

            if len(transcript_chunks) > 1:
                for idx, chunk_lines in enumerate(transcript_chunks[1:], start=2):
                    total_continuation_pages += 1
                    cont = out.new_page(width=page_width, height=page_height * 0.55)
                    cont.insert_text((x_margin, 20), f"원본 슬라이드 {logical_index + 1} 전사문 이어짐 ({idx})", fontsize=9.5)
                    cont_box = fitz.Rect(x_margin, 36, page_width - x_margin, page_height * 0.55 - 16)
                    cont.draw_rect(cont_box, color=(0.85, 0.85, 0.85), width=0.8)
                    cont_img = os.path.join(temp_img_dir, f"cont_{logical_index+1}_{idx}.png")
                    render_lines_panel_to_png(
                        chunk_lines,
                        int(cont_box.width),
                        int(cont_box.height),
                        fontfile,
                        used_tr_font,
                        cont_img,
                    )
                    cont.insert_image(cont_box, filename=cont_img)

        out.save(output_pdf_path)
        out.close()

        document_title = "전체 전사문 요약 정리표"
        subtitle = next((b.source_name for b in blocks if b.source_name.strip()), "줄+야붙 최종본 전사문")
        digest_info = render_transcript_digest_outputs(digest_entries, digest_pdf_path, fontfile, document_title, subtitle, output_dir)
        digest_report_path = os.path.join(output_dir, "overall_transcript_summary_digest_generation.json")
        with open(digest_report_path, "w", encoding="utf-8") as f:
            json.dump(digest_info, f, ensure_ascii=False, indent=2)

        coverage_review, coverage_txt = build_rendering_coverage_report(blocks, rendered_transcript_fragments, block_diagnostics)
        completeness_report = {
            "generated_pdf": os.path.basename(output_pdf_path),
            "source_pdf_map": source_pdf_map,
            "render_source_policy": render_mode_label,
            "block_assignment_summary": diagnostic_summary,
            "block_assignments": block_diagnostics,
            "page_attachment_summary": page_attachment_summary,
            "coverage_review": coverage_review,
            "continuation_pages_generated": total_continuation_pages,
        }
        with open(completeness_json_path, "w", encoding="utf-8") as f:
            json.dump(completeness_report, f, ensure_ascii=False, indent=2)
        with open(completeness_txt_path, "w", encoding="utf-8") as f:
            f.write(coverage_txt)

        original_transcript = "\n".join(f"[{b.timecode}] {b.body}" for b in blocks)
        finalcheck = {
            "generated_pdf": os.path.basename(output_pdf_path),
            "generated_transcript_digest_pdf": os.path.basename(digest_pdf_path),
            "generated_transcript_digest_docx": "overall_transcript_summary_digest.docx",
            "fontfile": fontfile,
            "garble_suspected": False,
            "transcript_token_coverage": 1.0 if coverage_review["all_blocks_rendered_once_in_order"] else 0.0,
            "removed_content_suspected": not coverage_review["all_blocks_rendered_once_in_order"],
            "all_source_pages_have_primary_output_page": True,
            "logical_source_pages": len(page_texts),
            "overlap_terms_by_logical_page": all_overlaps,
            "rendering_mode": render_mode_label,
            "transcript_completeness_review_json": os.path.basename(completeness_json_path),
            "review_needed": coverage_review["needs_manual_review"],
            "continuation_policy": "unlimited_pages_keep_all_transcript",
            "continuation_pages_generated": total_continuation_pages,
            "digest_entry_count": len(digest_entries),
            "digest_layout": "sectioned_portrait_summary_tables",
            "special_marker_label": SPECIAL_MARKER_LABEL,
            "special_marker_page_count": sum(1 for x in special_marker_pages if x),
            "special_marker_source_pages": [i + 1 for i, x in enumerate(special_marker_pages) if x],
            "red_emphasis_pages": [i + 1 for i, pf in page_features.items() if pf.get("red_fragments")],
            "original_transcript_char_count": len(original_transcript),
        }
        with open(finalcheck_path, "w", encoding="utf-8") as f:
            json.dump(finalcheck, f, ensure_ascii=False, indent=2)

        return output_pdf_path, precheck_path, finalcheck_path, completeness_json_path, completeness_txt_path
    finally:
        for d in docs.values():
            d.close()


def generate_annotated_pdf(pdf_path: str, transcript_text: str, output_dir: str) -> Tuple[str, str, str, str, str]:
    page_plan, source_pdf_map = _build_single_pdf_plan(pdf_path)
    return _generate_annotated_pdf_core(
        page_plan=page_plan,
        source_pdf_map=source_pdf_map,
        transcript_text=transcript_text,
        output_dir=output_dir,
        render_mode_label="single_pdf_direct_render_with_annots",
    )


def generate_annotated_pdf_from_plan(
    page_plan: List[Dict[str, Any]] | List[LogicalPageRef],
    source_pdf_map: Dict[str, str],
    transcript_text: str,
    output_dir: str,
) -> Tuple[str, str, str, str, str]:
    logical_plan: List[LogicalPageRef] = []
    for item in page_plan:
        if isinstance(item, LogicalPageRef):
            logical_plan.append(item)
        else:
            logical_plan.append(LogicalPageRef(source=item["source"], page_no=int(item["page_no"]), reason=item.get("reason")))
    return _generate_annotated_pdf_core(
        page_plan=logical_plan,
        source_pdf_map=source_pdf_map,
        transcript_text=transcript_text,
        output_dir=output_dir,
        render_mode_label="merged_plan_direct_render_from_original_sources_with_annots",
    )


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Generate annotated lecture note PDF from slide PDF and paged transcript.")
    parser.add_argument("pdf_path")
    parser.add_argument("transcript_path")
    parser.add_argument("--output", default="output")
    args = parser.parse_args()

    transcript_text = Path(args.transcript_path).read_text(encoding="utf-8")
    outputs = generate_annotated_pdf(args.pdf_path, transcript_text, args.output)
    for item in outputs:
        print(item)
